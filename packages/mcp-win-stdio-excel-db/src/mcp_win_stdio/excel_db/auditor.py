#!/usr/bin/env python3
"""
Automated Data Auditor and Reconciliation Engine for comparing Database tables
and Excel spreadsheets with tolerance, column mappings, and multi-tab diff reports.
"""

import math
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import create_engine, text

from mcp_win_stdio.excel_db.stream import resolve_sqlalchemy_url


def _clean_val(v: Any) -> Any:
    """Format primitives safely for JSON serialization and comparison."""
    if pd.isna(v):
        return None
    elif hasattr(v, "isoformat"):
        return v.isoformat()
    elif isinstance(v, (float, int)) and (math.isnan(v) or math.isinf(v)):
        return None
    return v


def _load_dataset(
    source: Union[str, Dict[str, Any]],
    db_resolver_func = None
) -> pd.DataFrame:
    """
    Load a dataset from either a file path (Excel/CSV) or a configuration dict (DB/Excel).
    """
    if isinstance(source, str):
        path = Path(source).resolve()
        if not path.exists():
            raise FileNotFoundError(f"Source file not found: {path}")
        if path.suffix.lower() == ".csv":
            df = pd.read_csv(path)
        else:
            df = pd.read_excel(path)

        if len(df) > 0 and (df.columns.dtype == 'int64' or all(str(c).isdigit() for c in df.columns)):
            new_cols = [str(x).strip() for x in df.iloc[0]]
            df = df.iloc[1:].reset_index(drop=True)
            df.columns = new_cols
        return df

    if not isinstance(source, dict):
        raise ValueError(f"Invalid source specification: {source}")

    s_type = source.get("type", "excel").lower()

    if s_type == "db":
        conn_id = source.get("connection")
        raw_sql = source.get("query")
        if not raw_sql:
            raise ValueError("Database source must specify 'query'.")
        db_url = db_resolver_func(conn_id) if db_resolver_func else conn_id
        resolved_url = resolve_sqlalchemy_url(db_url)
        engine = create_engine(resolved_url)
        with engine.connect() as conn:
            return pd.read_sql_query(text(raw_sql), conn)

    elif s_type in ("excel", "xlsx"):
        path = Path(source.get("path", "")).resolve()
        if not path.exists():
            raise FileNotFoundError(f"Source Excel file not found: {path}")
        sheet = source.get("sheet", 0)
        return pd.read_excel(path, sheet_name=sheet)

    elif s_type == "csv":
        path = Path(source.get("path", "")).resolve()
        if not path.exists():
            raise FileNotFoundError(f"Source CSV file not found: {path}")
        return pd.read_csv(path)

    else:
        raise ValueError(f"Unsupported source type: '{s_type}'.")


def compare_master_datasets(
    source_a: Union[str, Dict[str, Any]],
    source_b: Union[str, Dict[str, Any]],
    key_columns: List[str],
    column_mapping: Optional[Dict[str, str]] = None,
    compare_columns: Optional[List[str]] = None,
    numeric_tolerance: float = 0.001,
    ignore_whitespace_case: bool = True,
    output_report_path: Optional[str] = None,
    db_resolver_func = None
) -> Dict[str, Any]:
    """
    Deterministically compare two master datasets (Excel vs Excel, DB vs Excel, or DB vs DB).
    Detects new records, missing records, and field-level discrepancies with tolerance.

    Args:
        source_a: Baseline master (file path or dict specification).
        source_b: Comparison / New master (file path or dict specification).
        key_columns: List of primary key column names (as named in source_a).
        column_mapping: Optional map of source_b column names to source_a column names
                        (e.g. {"Material Number": "material_number", "Unit Rate": "sale_price"}).
        compare_columns: Optional subset of source_a column names to compare.
        numeric_tolerance: Maximum absolute difference considered equal for numbers (default 0.001).
        ignore_whitespace_case: If True, trims whitespace and ignores case for string comparison.
        output_report_path: Optional path to save a styled multi-tab Excel audit report.
        db_resolver_func: Function to resolve connection names if DB sources are used.
    """
    t_start = datetime.now()

    # 1. Load DataFrames
    df_a = _load_dataset(source_a, db_resolver_func)
    df_b = _load_dataset(source_b, db_resolver_func)

    # 2. Normalize and apply column mapping to source_b
    df_a.columns = [str(c).strip() for c in df_a.columns]
    df_b.columns = [str(c).strip() for c in df_b.columns]

    if column_mapping:
        # Rename source_b columns to match source_a naming convention
        rename_dict = {}
        for b_col, a_col in column_mapping.items():
            b_clean = str(b_col).strip()
            a_clean = str(a_col).strip()
            if b_clean in df_b.columns:
                rename_dict[b_clean] = a_clean
        df_b = df_b.rename(columns=rename_dict)

    # 3. Validate Key Columns
    for k in key_columns:
        if k not in df_a.columns:
            raise ValueError(f"Key column '{k}' not found in Source A. Available: {list(df_a.columns)}")
        if k not in df_b.columns:
            raise ValueError(f"Key column '{k}' not found in Source B. Available: {list(df_b.columns)}")

    # 4. Determine columns to compare
    if not compare_columns:
        shared_cols = [c for c in df_a.columns if c in df_b.columns and c not in key_columns]
    else:
        shared_cols = [c for c in compare_columns if c in df_a.columns and c in df_b.columns and c not in key_columns]

    # Clean key column values to strings for robust joining
    for k in key_columns:
        df_a[k] = df_a[k].astype(str).str.strip()
        df_b[k] = df_b[k].astype(str).str.strip()

    # 5. Full Outer Join on Key Columns
    merged = pd.merge(
        df_a, df_b,
        on=key_columns,
        how="outer",
        suffixes=("_a", "_b"),
        indicator=True
    )

    only_in_a = merged[merged["_merge"] == "left_only"].copy()
    only_in_b = merged[merged["_merge"] == "right_only"].copy()
    in_both = merged[merged["_merge"] == "both"].copy()

    # 6. Detailed Field Comparison
    discrepancies = []
    col_discrepancy_counts: Dict[str, int] = {c: 0 for c in shared_cols}

    for _, row in in_both.iterrows():
        row_diffs = {}
        for col in shared_cols:
            val_a = row.get(f"{col}_a")
            val_b = row.get(f"{col}_b")

            # Handle nulls
            is_na_a = pd.isna(val_a)
            is_na_b = pd.isna(val_b)
            if is_na_a and is_na_b:
                continue
            if is_na_a != is_na_b:
                row_diffs[col] = {
                    "source_a": _clean_val(val_a),
                    "source_b": _clean_val(val_b),
                    "delta": None,
                    "type": "null_mismatch"
                }
                col_discrepancy_counts[col] += 1
                continue

            # Numeric comparison
            is_num_a = isinstance(val_a, (int, float)) and not isinstance(val_a, bool)
            is_num_b = isinstance(val_b, (int, float)) and not isinstance(val_b, bool)

            if is_num_a and is_num_b:
                delta = float(val_b) - float(val_a)
                if abs(delta) > numeric_tolerance:
                    pct = round((delta / float(val_a) * 100), 2) if float(val_a) != 0 else None
                    row_diffs[col] = {
                        "source_a": _clean_val(val_a),
                        "source_b": _clean_val(val_b),
                        "delta": round(delta, 4),
                        "pct_change": pct,
                        "type": "numeric_diff"
                    }
                    col_discrepancy_counts[col] += 1
            else:
                # String / Boolean comparison
                s_a = str(val_a).strip()
                s_b = str(val_b).strip()
                if ignore_whitespace_case:
                    s_a_cmp = s_a.lower()
                    s_b_cmp = s_b.lower()
                else:
                    s_a_cmp = s_a
                    s_b_cmp = s_b

                if s_a_cmp != s_b_cmp:
                    row_diffs[col] = {
                        "source_a": _clean_val(val_a),
                        "source_b": _clean_val(val_b),
                        "delta": None,
                        "type": "text_diff"
                    }
                    col_discrepancy_counts[col] += 1

        if row_diffs:
            keys_dict = {k: row[k] for k in key_columns}
            discrepancies.append({
                "keys": keys_dict,
                "differences": row_diffs
            })

    identical_records_count = len(in_both) - len(discrepancies)

    # 7. Generate Multi-Tab Styled Audit Report
    report_file_path = None
    if output_report_path:
        out_p = Path(output_report_path).resolve()
        out_p.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()

        # Styles
        font_title = Font(name="Calibri", size=14, bold=True, color="1F497D")
        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        font_bold = Font(name="Calibri", size=11, bold=True)
        fill_header_navy = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        fill_header_green = PatternFill(start_color="27AE60", end_color="27AE60", fill_type="solid")
        fill_header_orange = PatternFill(start_color="E67E22", end_color="E67E22", fill_type="solid")
        fill_diff_yellow = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
        fill_diff_red = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
        border_thin = Border(
            left=Side(style="thin", color="D3D3D3"),
            right=Side(style="thin", color="D3D3D3"),
            top=Side(style="thin", color="D3D3D3"),
            bottom=Side(style="thin", color="D3D3D3")
        )

        # Tab 1: Overview
        ws_overview = wb.active
        ws_overview.title = "Audit Overview"
        ws_overview.append(["Master Data Reconciliation Audit Report", ""])
        ws_overview["A1"].font = font_title
        ws_overview.append([])
        ws_overview.append(["Metric", "Count / Value"])
        ws_overview["A3"].font = font_header
        ws_overview["B3"].font = font_header
        ws_overview["A3"].fill = fill_header_navy
        ws_overview["B3"].fill = fill_header_navy

        metrics = [
            ("Audit Execution Timestamp", datetime.now().isoformat()),
            ("Source A Total Records", len(df_a)),
            ("Source B Total Records", len(df_b)),
            ("Matched Primary Keys", len(in_both)),
            ("Identical Records (Within Tolerance)", identical_records_count),
            ("Records with Value Discrepancies", len(discrepancies)),
            ("New Records (Source B Only)", len(only_in_b)),
            ("Missing Records (Source A Only)", len(only_in_a)),
            ("Numeric Tolerance Applied", numeric_tolerance),
        ]
        for m_name, m_val in metrics:
            ws_overview.append([m_name, m_val])

        ws_overview.append([])
        ws_overview.append(["Field Name", "Mismatched Rows Count"])
        ws_overview["A15"].font = font_header
        ws_overview["B15"].font = font_header
        ws_overview["A15"].fill = fill_header_navy
        ws_overview["B15"].fill = fill_header_navy
        for fld, cnt in col_discrepancy_counts.items():
            ws_overview.append([fld, cnt])

        # Tab 2: Discrepancies
        if discrepancies:
            ws_diff = wb.create_sheet("Field Differences")
            diff_headers = list(key_columns) + ["Field", "Source A (Baseline)", "Source B (New)", "Delta", "% Change", "Diff Type"]
            ws_diff.append(diff_headers)
            for c_idx in range(1, len(diff_headers) + 1):
                cell = ws_diff.cell(row=1, column=c_idx)
                cell.font = font_header
                cell.fill = fill_header_navy

            for item in discrepancies:
                k_vals = [item["keys"][k] for k in key_columns]
                for fld, diff_info in item["differences"].items():
                    row_vals = k_vals + [
                        fld,
                        str(diff_info["source_a"]),
                        str(diff_info["source_b"]),
                        diff_info.get("delta"),
                        diff_info.get("pct_change"),
                        diff_info.get("type")
                    ]
                    ws_diff.append(row_vals)
                    # highlight
                    curr_row = ws_diff.max_row
                    for c_idx in range(1, len(row_vals) + 1):
                        ws_diff.cell(row=curr_row, column=c_idx).fill = fill_diff_yellow

        # Tab 3: New Records in B
        if len(only_in_b) > 0:
            ws_new = wb.create_sheet("New Records")
            b_cols = [c for c in df_b.columns if not c.endswith("_a")]
            clean_b_cols = [c.replace("_b", "") for c in b_cols]
            ws_new.append(clean_b_cols)
            for c_idx in range(1, len(clean_b_cols) + 1):
                cell = ws_new.cell(row=1, column=c_idx)
                cell.font = font_header
                cell.fill = fill_header_green

            for _, row in only_in_b.iterrows():
                ws_new.append([_clean_val(row.get(f"{c}_b", row.get(c))) for c in clean_b_cols])

        # Tab 4: Missing Records in B
        if len(only_in_a) > 0:
            ws_miss = wb.create_sheet("Missing Records")
            a_cols = [c for c in df_a.columns if not c.endswith("_b")]
            clean_a_cols = [c.replace("_a", "") for c in a_cols]
            ws_miss.append(clean_a_cols)
            for c_idx in range(1, len(clean_a_cols) + 1):
                cell = ws_miss.cell(row=1, column=c_idx)
                cell.font = font_header
                cell.fill = fill_header_orange

            for _, row in only_in_a.iterrows():
                ws_miss.append([_clean_val(row.get(f"{c}_a", row.get(c))) for c in clean_a_cols])

        # Auto-fit columns across all sheets
        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_len = max(len(str(c.value or "")) for c in col)
                col_letter = get_column_letter(col[0].column)
                sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        wb.save(out_p)
        wb.close()
        report_file_path = str(out_p)

    duration = (datetime.now() - t_start).total_seconds()

    return {
        "status": "success",
        "summary": {
            "source_a_records": len(df_a),
            "source_b_records": len(df_b),
            "matched_keys": len(in_both),
            "matching_records": identical_records_count,
            "modified_records": len(discrepancies),
            "new_records_in_b": len(only_in_b),
            "missing_records_in_b": len(only_in_a),
        },
        "total_source_a_records": len(df_a),
        "total_source_b_records": len(df_b),
        "matched_primary_keys": len(in_both),
        "identical_records_count": identical_records_count,
        "records_with_discrepancies_count": len(discrepancies),
        "new_records_in_source_b_count": len(only_in_b),
        "missing_records_in_source_b_count": len(only_in_a),
        "column_discrepancy_counts": col_discrepancy_counts,
        "sample_discrepancies": discrepancies[:10],
        "audit_report_file": report_file_path,
        "duration_seconds": round(duration, 3)
    }


def reconcile_db_vs_excel(
    sql_query: str,
    excel_path: str,
    key_columns: List[str],
    db_url_or_config: Union[str, Dict[str, Any]],
    sheet_name: Optional[Union[str, int]] = 0,
    compare_columns: Optional[List[str]] = None,
    output_report_path: Optional[str] = None
) -> Dict[str, Any]:
    """
    Backwards-compatible wrapper delegating to compare_master_datasets.
    """
    db_source = {
        "type": "db",
        "connection": db_url_or_config,
        "query": sql_query
    }
    excel_source = {
        "type": "excel",
        "path": excel_path,
        "sheet": sheet_name
    }
    return compare_master_datasets(
        source_a=db_source,
        source_b=excel_source,
        key_columns=key_columns,
        compare_columns=compare_columns,
        output_report_path=output_report_path,
        db_resolver_func=lambda c: c
    )
