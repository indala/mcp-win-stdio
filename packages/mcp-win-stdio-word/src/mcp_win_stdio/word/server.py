#!/usr/bin/env python3
"""
Word Document MCP Server for Windows (Advanced Layout & Typography Edition)
Extracts comprehensive content, headers, footers, advanced page geometry, exact margins across all units,
multi-column layouts (IEEE/Journal), paragraph spacing/indentation, typography, tables, outline, images, and metadata.
"""

import json
import os
import re
from typing import Any, Dict, List, Optional

try:
    from mcp.server.mcpserver import MCPServer as FastMCP
except (ImportError, ModuleNotFoundError):
    from mcp.server.fastmcp import FastMCP
import docx
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import nsmap
from docx.shared import Inches, Pt, RGBColor

# Register namespaces
nsmap["wp"] = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
nsmap["a"] = "http://schemas.openxmlformats.org/drawingml/2006/main"
nsmap["pic"] = "http://schemas.openxmlformats.org/drawingml/2006/picture"
nsmap["r"] = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
nsmap["v"] = "urn:schemas-microsoft-com:vml"

mcp = FastMCP("word-mcp")

# Unit Conversion Constants
EMU_PER_INCH = 914400.0
EMU_PER_CM = 360000.0
EMU_PER_PT = 12700.0
TWIPS_PER_INCH = 1440.0
TWIPS_PER_CM = 567.0
TWIPS_PER_PT = 20.0


def identify_paper_size(width_in: float, height_in: float) -> str:
    w, h = sorted([round(width_in, 2), round(height_in, 2)])
    if abs(w - 8.27) < 0.05 and abs(h - 11.69) < 0.05:
        return "A4 (8.27 x 11.69 in / 210 x 297 mm)"
    elif abs(w - 8.5) < 0.05 and abs(h - 11.0) < 0.05:
        return "Letter (8.5 x 11 in / 216 x 279 mm)"
    elif abs(w - 8.5) < 0.05 and abs(h - 14.0) < 0.05:
        return "Legal (8.5 x 14 in / 216 x 356 mm)"
    elif abs(w - 5.83) < 0.05 and abs(h - 8.27) < 0.05:
        return "A5 (5.83 x 8.27 in / 148 x 210 mm)"
    elif abs(w - 11.69) < 0.05 and abs(h - 16.54) < 0.05:
        return "A3 (11.69 x 16.54 in / 297 x 420 mm)"
    elif abs(w - 7.25) < 0.05 and abs(h - 10.5) < 0.05:
        return "Executive (7.25 x 10.5 in / 184 x 267 mm)"
    elif abs(w - 7.17) < 0.05 and abs(h - 10.12) < 0.05:
        return "B5 JIS (7.17 x 10.12 in / 182 x 257 mm)"
    return f"Custom ({width_in:.2f} x {height_in:.2f} in)"


def load_document(file_path: str) -> docx.Document:
    clean_path = os.path.abspath(file_path.strip("\"'"))
    if not os.path.exists(clean_path):
        raise FileNotFoundError(f"File not found: {clean_path}")

    if clean_path.lower().endswith(".doc") and not clean_path.lower().endswith(".docx"):
        try:
            import pythoncom
            import win32com.client

            pythoncom.CoInitialize()
            word = None
            doc = None
            try:
                word = win32com.client.Dispatch("Word.Application")
                word.Visible = False
                word.DisplayAlerts = False
                doc = word.Documents.Open(clean_path)
                temp_docx = clean_path + "x"
                doc.SaveAs2(temp_docx, FileFormat=16)
                doc.Close(False)
                doc = None
                return docx.Document(temp_docx)
            finally:
                if doc:
                    try:
                        doc.Close(False)
                    except Exception:
                        pass
                if word:
                    try:
                        word.Quit()
                    except Exception:
                        pass
                pythoncom.CoUninitialize()
        except Exception as e:
            raise RuntimeError(f"Could not convert legacy .doc file using MS Word: {e}")

    return docx.Document(clean_path)


def extract_run_formatting(run, paragraph_style=None) -> Dict[str, Any]:
    """Extract font name, size (pt), color (hex/theme), highlight, and bold/italic."""
    font_name = run.font.name
    if not font_name and paragraph_style and hasattr(paragraph_style, "font"):
        font_name = paragraph_style.font.name

    font_size = None
    if run.font.size:
        font_size = run.font.size.pt
    elif paragraph_style and hasattr(paragraph_style, "font") and paragraph_style.font.size:
        font_size = paragraph_style.font.size.pt

    color_hex = None
    if run.font.color and run.font.color.rgb:
        color_hex = f"#{run.font.color.rgb}"

    highlight = str(run.font.highlight_color) if run.font.highlight_color else None

    rPr = run._r.xpath('.//*[local-name()="rPr"]')
    if rPr:
        if not color_hex:
            color_el = rPr[0].xpath('.//*[local-name()="color"]')
            if color_el:
                val = color_el[0].get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") or color_el[
                    0
                ].get("val")
                if val and val != "auto":
                    color_hex = f"#{val}"
                theme = color_el[0].get(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}themeColor"
                ) or color_el[0].get("themeColor")
                if theme and not color_hex:
                    color_hex = f"theme:{theme}"
        if not font_size:
            sz_el = rPr[0].xpath('.//*[local-name()="sz"]')
            if sz_el:
                val = sz_el[0].get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") or sz_el[0].get(
                    "val"
                )
                if val:
                    try:
                        font_size = int(val) / 2.0
                    except:
                        pass
        if not font_name:
            rFonts_el = rPr[0].xpath('.//*[local-name()="rFonts"]')
            if rFonts_el:
                font_name = rFonts_el[0].get(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii"
                ) or rFonts_el[0].get("ascii")

    return {
        "font_family": font_name or "Default / Body Font",
        "font_size_pt": font_size if font_size is not None else "Default / 11pt",
        "font_color": color_hex or "#000000 (Automatic)",
        "highlight": highlight,
        "bold": bool(run.bold),
        "italic": bool(run.italic),
        "underline": bool(run.underline),
    }


def extract_drawing_info(element, location: str, part=None) -> List[Dict[str, Any]]:
    images = []
    if element is None:
        return images

    drawings = element.xpath('.//*[local-name()="drawing"]')

    for d in drawings:
        anchors = d.xpath('.//*[local-name()="anchor"]')
        inlines = d.xpath('.//*[local-name()="inline"]')

        for item in anchors + inlines:
            is_anchor = item in anchors
            placement_type = "FLOATING_ANCHOR" if is_anchor else "INLINE"

            ext = item.xpath('.//*[local-name()="extent"]')
            cx = int(ext[0].get("cx", 0)) if ext else 0
            cy = int(ext[0].get("cy", 0)) if ext else 0

            doc_pr = item.xpath('.//*[local-name()="docPr"]')
            name = doc_pr[0].get("name", "") if doc_pr else ""
            descr = doc_pr[0].get("descr", "") if doc_pr else ""
            title = doc_pr[0].get("title", "") if doc_pr else ""

            blips = item.xpath('.//*[local-name()="blip"]')
            r_id = ""
            if blips:
                for attr_name, attr_val in blips[0].attrib.items():
                    if attr_name.endswith("embed"):
                        r_id = attr_val
                        break

            target_file = ""
            content_type = ""
            file_size_bytes = None

            if r_id and part and hasattr(part, "rels"):
                try:
                    rel = part.rels.get(r_id)
                    if rel:
                        target_file = os.path.basename(rel.target_ref)
                        if hasattr(part, "related_parts"):
                            related_part = part.related_parts.get(r_id)
                            if related_part:
                                content_type = getattr(related_part, "content_type", "")
                                blob = getattr(related_part, "blob", None)
                                if blob:
                                    file_size_bytes = len(blob)
                except Exception:
                    pass

            h_align = None
            h_rel_from = None
            h_offset_in = None
            pos_h = item.xpath('.//*[local-name()="positionH"]')
            if pos_h:
                h_rel_from = pos_h[0].get("relativeFrom", "")
                align_el = pos_h[0].xpath('.//*[local-name()="align"]')
                if align_el and align_el[0].text:
                    h_align = align_el[0].text.strip()
                pos_offset = pos_h[0].xpath('.//*[local-name()="posOffset"]')
                if pos_offset and pos_offset[0].text:
                    try:
                        h_offset_in = round(int(pos_offset[0].text) / EMU_PER_INCH, 2)
                    except:
                        pass

            v_align = None
            v_rel_from = None
            v_offset_in = None
            pos_v = item.xpath('.//*[local-name()="positionV"]')
            if pos_v:
                v_rel_from = pos_v[0].get("relativeFrom", "")
                align_el = pos_v[0].xpath('.//*[local-name()="align"]')
                if align_el and align_el[0].text:
                    v_align = align_el[0].text.strip()
                pos_offset = pos_v[0].xpath('.//*[local-name()="posOffset"]')
                if pos_offset and pos_offset[0].text:
                    try:
                        v_offset_in = round(int(pos_offset[0].text) / EMU_PER_INCH, 2)
                    except:
                        pass

            images.append(
                {
                    "location": location,
                    "placement_type": placement_type,
                    "image_name": name or "Unnamed Image",
                    "alt_text": descr or title or None,
                    "dimensions": {
                        "width_inches": round(cx / EMU_PER_INCH, 2),
                        "height_inches": round(cy / EMU_PER_INCH, 2),
                        "width_cm": round(cx / EMU_PER_CM, 2),
                        "height_cm": round(cy / EMU_PER_CM, 2),
                        "width_pt": round(cx / EMU_PER_PT, 1),
                        "height_pt": round(cy / EMU_PER_PT, 1),
                    },
                    "positioning": {
                        "horizontal_alignment": h_align or ("offset" if h_offset_in is not None else "inline"),
                        "horizontal_relative_from": h_rel_from or None,
                        "horizontal_offset_inches": h_offset_in,
                        "vertical_alignment": v_align or ("offset" if v_offset_in is not None else "inline"),
                        "vertical_relative_from": v_rel_from or None,
                        "vertical_offset_inches": v_offset_in,
                    },
                    "embedded_file": {
                        "target_filename": target_file or None,
                        "content_type": content_type or None,
                        "file_size_bytes": file_size_bytes,
                    },
                }
            )

    return images


@mcp.tool()
def get_document_layout(file_path: str) -> Dict[str, Any]:
    """
    Extract comprehensive page layout geometry, exact margins across multiple units (in, cm, mm, pt, twips),
    multi-column layout (IEEE/Journal), printable area, section break types, and header/footer distances.

    Args:
        file_path: Absolute or relative path to the .docx or .doc file.
    """
    try:
        doc = load_document(file_path)
        sections_info = []

        for idx, section in enumerate(doc.sections, 1):
            sectPr = section._sectPr

            # Dimensions
            w_in = section.page_width.inches if section.page_width else 8.5
            h_in = section.page_height.inches if section.page_height else 11.0
            orientation_str = "PORTRAIT" if section.orientation == WD_ORIENT.PORTRAIT else "LANDSCAPE"
            paper_size = identify_paper_size(w_in, h_in)

            # Margins
            top_in = section.top_margin.inches if section.top_margin else 1.0
            bot_in = section.bottom_margin.inches if section.bottom_margin else 1.0
            left_in = section.left_margin.inches if section.left_margin else 1.0
            right_in = section.right_margin.inches if section.right_margin else 1.0
            gutter_in = section.gutter.inches if section.gutter else 0.0
            header_dist_in = section.header_distance.inches if section.header_distance else 0.5
            footer_dist_in = section.footer_distance.inches if section.footer_distance else 0.5

            # Printable Area
            printable_w_in = round(w_in - left_in - right_in - gutter_in, 2)
            printable_h_in = round(h_in - top_in - bot_in, 2)

            # Section Type & Flow
            type_el = sectPr.xpath('.//*[local-name()="type"]')
            sec_type_raw = (
                type_el[0].get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")
                or type_el[0].get("val")
                if type_el
                else "nextPage"
            )
            sec_type_map = {
                "nextPage": "NEW_PAGE (Starts on Next Page)",
                "continuous": "CONTINUOUS (Continuous Section Break)",
                "evenPage": "EVEN_PAGE (Starts on Next Even Page)",
                "oddPage": "ODD_PAGE (Starts on Next Odd Page)",
                "nextColumn": "NEXT_COLUMN (Starts on Next Column)",
            }
            sec_type_readable = sec_type_map.get(sec_type_raw, sec_type_raw)

            # Multi-Column Layout (IEEE / Journal 2-Column layout)
            cols_el = sectPr.xpath('.//*[local-name()="cols"]')
            num_cols = 1
            col_space_in = None
            equal_width = True
            line_sep = False
            col_width_in = printable_w_in

            if cols_el:
                num_attr = cols_el[0].get(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num"
                ) or cols_el[0].get("num")
                if num_attr:
                    try:
                        num_cols = int(num_attr)
                    except:
                        num_cols = 1
                space_attr = cols_el[0].get(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}space"
                ) or cols_el[0].get("space")
                if space_attr:
                    try:
                        col_space_in = round(int(space_attr) / TWIPS_PER_INCH, 2)
                    except:
                        pass
                sep_attr = cols_el[0].get(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sep"
                ) or cols_el[0].get("sep")
                line_sep = bool(sep_attr and sep_attr in ["1", "true", "on"])
                eq_attr = cols_el[0].get(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}equalWidth"
                ) or cols_el[0].get("equalWidth")
                equal_width = eq_attr not in ["0", "false", "off"]

                if num_cols > 1:
                    spacing_total = (col_space_in or 0.25) * (num_cols - 1)
                    col_width_in = round((printable_w_in - spacing_total) / num_cols, 2)

            # Page Numbering
            pgNum_el = sectPr.xpath('.//*[local-name()="pgNumType"]')
            pg_start = None
            pg_format = "decimal (1, 2, 3...)"
            if pgNum_el:
                start_attr = pgNum_el[0].get(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}start"
                ) or pgNum_el[0].get("start")
                if start_attr:
                    try:
                        pg_start = int(start_attr)
                    except:
                        pass
                fmt_attr = pgNum_el[0].get(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fmt"
                ) or pgNum_el[0].get("fmt")
                if fmt_attr:
                    pg_format = fmt_attr

            sections_info.append(
                {
                    "section_number": idx,
                    "section_break_type": sec_type_readable,
                    "orientation": orientation_str,
                    "paper_size": paper_size,
                    "page_dimensions": {
                        "inches": {"width": round(w_in, 2), "height": round(h_in, 2)},
                        "centimeters": {"width": round(w_in * 2.54, 2), "height": round(h_in * 2.54, 2)},
                        "millimeters": {"width": round(w_in * 25.4, 1), "height": round(h_in * 25.4, 1)},
                        "points": {"width": round(w_in * 72.0, 1), "height": round(h_in * 72.0, 1)},
                    },
                    "printable_area": {
                        "inches": {"width": printable_w_in, "height": printable_h_in},
                        "centimeters": {
                            "width": round(printable_w_in * 2.54, 2),
                            "height": round(printable_h_in * 2.54, 2),
                        },
                        "points": {"width": round(printable_w_in * 72.0, 1), "height": round(printable_h_in * 72.0, 1)},
                    },
                    "margins": {
                        "inches": {
                            "top": round(top_in, 2),
                            "bottom": round(bot_in, 2),
                            "left": round(left_in, 2),
                            "right": round(right_in, 2),
                            "gutter": round(gutter_in, 2),
                        },
                        "centimeters": {
                            "top": round(top_in * 2.54, 2),
                            "bottom": round(bot_in * 2.54, 2),
                            "left": round(left_in * 2.54, 2),
                            "right": round(right_in * 2.54, 2),
                            "gutter": round(gutter_in * 2.54, 2),
                        },
                        "millimeters": {
                            "top": round(top_in * 25.4, 1),
                            "bottom": round(bot_in * 25.4, 1),
                            "left": round(left_in * 25.4, 1),
                            "right": round(right_in * 25.4, 1),
                            "gutter": round(gutter_in * 25.4, 1),
                        },
                        "points": {
                            "top": round(top_in * 72.0, 1),
                            "bottom": round(bot_in * 72.0, 1),
                            "left": round(left_in * 72.0, 1),
                            "right": round(right_in * 72.0, 1),
                            "gutter": round(gutter_in * 72.0, 1),
                        },
                    },
                    "multi_column_layout": {
                        "is_multi_column": num_cols > 1,
                        "column_count": num_cols,
                        "column_width_inches": col_width_in,
                        "column_width_cm": round(col_width_in * 2.54, 2),
                        "column_spacing_inches": col_space_in or (0.25 if num_cols > 1 else 0.0),
                        "column_spacing_cm": round((col_space_in or (0.25 if num_cols > 1 else 0.0)) * 2.54, 2),
                        "equal_width_columns": equal_width,
                        "line_separator_between_columns": line_sep,
                    },
                    "header_footer_geometry": {
                        "header_distance_from_top_edge_inches": round(header_dist_in, 2),
                        "header_distance_from_top_edge_cm": round(header_dist_in * 2.54, 2),
                        "footer_distance_from_bottom_edge_inches": round(footer_dist_in, 2),
                        "footer_distance_from_bottom_edge_cm": round(footer_dist_in * 2.54, 2),
                        "different_first_page": section.different_first_page_header_footer,
                        "different_odd_and_even_pages": getattr(doc.settings, "odd_and_even_pages_header_footer", False)
                        if hasattr(doc, "settings")
                        else False,
                    },
                    "page_numbering": {
                        "starts_at": pg_start or "Continue from previous section",
                        "number_format": pg_format,
                    },
                }
            )

        return {"file_path": os.path.abspath(file_path), "total_sections": len(doc.sections), "sections": sections_info}
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def get_paragraph_spacing_and_indentation(file_path: str, max_paragraphs: int = 40) -> Dict[str, Any]:
    """
    Extract detailed paragraph spacing (line spacing, space before/after in pt), indentation (first-line indent, left/right margins in inches), and alignment.

    Args:
        file_path: Absolute or relative path to the .docx or .doc file.
        max_paragraphs: Number of paragraphs to analyze (default: 40, max: 100).
    """
    try:
        doc = load_document(file_path)
        safe_max = min(max(1, max_paragraphs), 100)
        paragraphs_info = []

        for idx, p in enumerate(doc.paragraphs[:safe_max]):
            text = p.text.strip()
            if not text:
                continue

            pf = p.paragraph_format
            align = str(p.alignment).split(".")[-1] if p.alignment else "LEFT"

            # Line Spacing
            line_spacing = pf.line_spacing
            line_spacing_rule = str(pf.line_spacing_rule).split(".")[-1] if pf.line_spacing_rule else "MULTIPLE"
            if line_spacing is None:
                line_spacing_desc = "Single (1.0 / Default)"
            elif isinstance(line_spacing, float):
                line_spacing_desc = f"{line_spacing:.2f}x Line Spacing"
            elif hasattr(line_spacing, "pt"):
                line_spacing_desc = f"Exact {line_spacing.pt:.1f} pt"
            else:
                line_spacing_desc = str(line_spacing)

            # Space Before & After
            space_before_pt = pf.space_before.pt if pf.space_before else 0.0
            space_after_pt = pf.space_after.pt if pf.space_after else 0.0

            # Indentation
            first_line_in = round(pf.first_line_indent.inches, 2) if pf.first_line_indent else 0.0
            left_indent_in = round(pf.left_indent.inches, 2) if pf.left_indent else 0.0
            right_indent_in = round(pf.right_indent.inches, 2) if pf.right_indent else 0.0

            paragraphs_info.append(
                {
                    "paragraph_index": idx + 1,
                    "style": p.style.name if p.style else "Normal",
                    "alignment": align,
                    "text_preview": text[:70] + ("..." if len(text) > 70 else ""),
                    "spacing": {
                        "line_spacing": line_spacing_desc,
                        "line_spacing_rule": line_spacing_rule,
                        "space_before_pt": space_before_pt,
                        "space_after_pt": space_after_pt,
                    },
                    "indentation": {
                        "first_line_indent_inches": first_line_in,
                        "first_line_indent_cm": round(first_line_in * 2.54, 2),
                        "left_indent_inches": left_indent_in,
                        "left_indent_cm": round(left_indent_in * 2.54, 2),
                        "right_indent_inches": right_indent_in,
                        "is_hanging_indent": first_line_in < 0,
                    },
                    "pagination_controls": {
                        "keep_with_next": bool(pf.keep_with_next),
                        "page_break_before": bool(pf.page_break_before),
                        "widow_control": bool(pf.widow_control),
                    },
                }
            )

        has_more = len(doc.paragraphs) > safe_max
        res: Dict[str, Any] = {
            "file_path": os.path.abspath(file_path),
            "total_paragraphs": len(doc.paragraphs),
            "analyzed_paragraphs_count": len(paragraphs_info),
            "has_more": has_more,
            "paragraphs": paragraphs_info,
        }
        if has_more:
            res["notice"] = (
                f"... [TRUNCATED: Showing first {safe_max} paragraphs. Increase max_paragraphs (up to 100) to see more] ..."
            )
        return res
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def get_document_typography(file_path: str, max_paragraphs: int = 40) -> Dict[str, Any]:
    """
    Extract all typography details (font families, font sizes, colors, headings, styles, and alignments) across the Word document.

    Args:
        file_path: Path to the .docx or .doc file.
        max_paragraphs: Maximum paragraph styles to detail (default: 40, max: 100).
    """
    try:
        doc = load_document(file_path)
        safe_max = min(max(1, max_paragraphs), 100)

        distinct_fonts = set()
        distinct_sizes = set()
        distinct_colors = set()
        styles_summary = []

        for p_idx, p in enumerate(doc.paragraphs):
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name if p.style else "Normal"
            align = str(p.alignment).split(".")[-1] if p.alignment else "LEFT"

            runs_data = []
            for r in p.runs:
                if r.text:
                    fmt = extract_run_formatting(r, p.style)
                    distinct_fonts.add(fmt["font_family"])
                    distinct_sizes.add(str(fmt["font_size_pt"]))
                    distinct_colors.add(fmt["font_color"])
                    runs_data.append({"text": r.text, "formatting": fmt})

            if runs_data and len(styles_summary) < safe_max:
                styles_summary.append(
                    {
                        "paragraph_index": p_idx + 1,
                        "style": style_name,
                        "alignment": align,
                        "text_preview": text[:80] + ("..." if len(text) > 80 else ""),
                        "runs": runs_data,
                    }
                )

        headers_typography = []
        for s_idx, sec in enumerate(doc.sections, 1):
            if sec.header:
                for p_idx, p in enumerate(sec.header.paragraphs):
                    if p.text.strip():
                        runs = [extract_run_formatting(r, p.style) for r in p.runs if r.text.strip()]
                        for r_fmt in runs:
                            distinct_fonts.add(r_fmt["font_family"])
                            distinct_sizes.add(str(r_fmt["font_size_pt"]))
                            distinct_colors.add(r_fmt["font_color"])
                        headers_typography.append(
                            {
                                "location": f"Section {s_idx} Header P#{p_idx + 1}",
                                "text": p.text.strip(),
                                "formatting": runs,
                            }
                        )

        return {
            "file_path": os.path.abspath(file_path),
            "summary": {
                "font_families_used": sorted(list(distinct_fonts)),
                "font_sizes_used_pt": sorted(list(distinct_sizes)),
                "colors_used": sorted(list(distinct_colors)),
            },
            "header_typography": headers_typography,
            "paragraphs_typography": styles_summary,
            "has_more": len(doc.paragraphs) > safe_max,
        }
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def get_document_images(file_path: str, max_images: int = 50) -> Dict[str, Any]:
    """
    Extract all images and graphics from headers, footers, body paragraphs, and tables along with exact placement, alignment, dimensions, and coordinates.

    Args:
        file_path: Absolute or relative path to the .docx or .doc file.
        max_images: Maximum images to detail (default: 50, max: 100).
    """
    try:
        doc = load_document(file_path)
        safe_max = min(max(1, max_images), 100)
        all_images = []

        for s_idx, sec in enumerate(doc.sections, 1):
            if sec.header:
                all_images.extend(
                    extract_drawing_info(sec.header._element, f"Section {s_idx} Primary Header", sec.header.part)
                )
            if sec.footer:
                all_images.extend(
                    extract_drawing_info(sec.footer._element, f"Section {s_idx} Primary Footer", sec.footer.part)
                )

            if sec.different_first_page_header_footer:
                if sec.first_page_header:
                    all_images.extend(
                        extract_drawing_info(
                            sec.first_page_header._element,
                            f"Section {s_idx} First Page Header",
                            sec.first_page_header.part,
                        )
                    )
                if sec.first_page_footer:
                    all_images.extend(
                        extract_drawing_info(
                            sec.first_page_footer._element,
                            f"Section {s_idx} First Page Footer",
                            sec.first_page_footer.part,
                        )
                    )

            if getattr(sec, "even_page_header", None):
                all_images.extend(
                    extract_drawing_info(
                        sec.even_page_header._element, f"Section {s_idx} Even Page Header", sec.even_page_header.part
                    )
                )
            if getattr(sec, "even_page_footer", None):
                all_images.extend(
                    extract_drawing_info(
                        sec.even_page_footer._element, f"Section {s_idx} Even Page Footer", sec.even_page_footer.part
                    )
                )

        for p_idx, p in enumerate(doc.paragraphs):
            p_images = extract_drawing_info(p._element, f"Body Paragraph #{p_idx + 1}", doc.part)
            all_images.extend(p_images)

        for t_idx, table in enumerate(doc.tables, 1):
            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    cell_images = extract_drawing_info(
                        cell._tc, f"Table {t_idx}, Cell ({r_idx + 1}, {c_idx + 1})", doc.part
                    )
                    all_images.extend(cell_images)

        has_more = len(all_images) > safe_max
        displayed_images = all_images[:safe_max]

        res: Dict[str, Any] = {
            "file_path": os.path.abspath(file_path),
            "total_images_found": len(all_images),
            "returned_images_count": len(displayed_images),
            "header_images_count": sum(1 for img in all_images if "Header" in img["location"]),
            "footer_images_count": sum(1 for img in all_images if "Footer" in img["location"]),
            "body_images_count": sum(
                1 for img in all_images if "Body" in img["location"] or "Table" in img["location"]
            ),
            "has_more": has_more,
            "images": displayed_images,
        }
        if has_more:
            res["notice"] = (
                f"... [TRUNCATED: Showing first {safe_max} images. Increase max_images (up to 100) to see more] ..."
            )
        return res
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def get_headers_and_footers(file_path: str) -> Dict[str, Any]:
    """
    Extract all headers and footers from each section of a Word document, including text, typography (fonts, sizes, colors), tables, and images.

    Args:
        file_path: Absolute or relative path to the .docx or .doc file.
    """
    try:
        doc = load_document(file_path)
        sections_hf = []

        for idx, section in enumerate(doc.sections, 1):

            def extract_hf_data(hf_obj, loc_name):
                if not hf_obj:
                    return None

                paras_info = []
                for p in hf_obj.paragraphs:
                    if p.text.strip():
                        runs = [
                            {"text": r.text, "formatting": extract_run_formatting(r, p.style)}
                            for r in p.runs
                            if r.text.strip()
                        ]
                        paras_info.append(
                            {
                                "text": p.text.strip(),
                                "alignment": str(p.alignment).split(".")[-1] if p.alignment else "LEFT",
                                "runs": runs,
                            }
                        )

                images = extract_drawing_info(hf_obj._element, loc_name, hf_obj.part)
                return {"paragraphs": paras_info, "image_count": len(images), "images": images}

            primary_h = extract_hf_data(section.header, f"Section {idx} Primary Header")
            primary_f = extract_hf_data(section.footer, f"Section {idx} Primary Footer")

            first_h = (
                extract_hf_data(section.first_page_header, f"Section {idx} First Page Header")
                if section.different_first_page_header_footer
                else None
            )
            first_f = (
                extract_hf_data(section.first_page_footer, f"Section {idx} First Page Footer")
                if section.different_first_page_header_footer
                else None
            )

            even_h = (
                extract_hf_data(getattr(section, "even_page_header", None), f"Section {idx} Even Page Header")
                if getattr(section, "even_page_header", None)
                else None
            )
            even_f = (
                extract_hf_data(getattr(section, "even_page_footer", None), f"Section {idx} Even Page Footer")
                if getattr(section, "even_page_footer", None)
                else None
            )

            sections_hf.append(
                {
                    "section_number": idx,
                    "different_first_page": section.different_first_page_header_footer,
                    "primary_header": primary_h,
                    "primary_footer": primary_f,
                    "first_page_header": first_h,
                    "first_page_footer": first_f,
                    "even_page_header": even_h,
                    "even_page_footer": even_f,
                }
            )

        return {"file_path": os.path.abspath(file_path), "sections": sections_hf}
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def get_document_tables(file_path: str, format: str = "markdown", max_tables: int = 10, max_rows: int = 30) -> Any:
    """
    Extract tables from a Word document as structured JSON or readable Markdown with safe limits.

    Args:
        file_path: Absolute or relative path to the .docx or .doc file.
        format: Output format ('markdown' or 'json'). Default is 'markdown'.
        max_tables: Maximum tables to extract (default: 10, max: 25).
        max_rows: Maximum rows per table to extract (default: 30, max: 100).
    """
    try:
        doc = load_document(file_path)
        tables_data = []
        safe_max_tables = min(max(1, max_tables), 25)
        safe_max_rows = min(max(1, max_rows), 100)

        for t_idx, table in enumerate(doc.tables[:safe_max_tables], 1):
            table_rows = []
            for row in table.rows[:safe_max_rows]:
                row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                table_rows.append(row_cells)

            tables_data.append(
                {
                    "table_index": t_idx,
                    "total_rows": len(table.rows),
                    "returned_rows": len(table_rows),
                    "total_columns": len(table.columns),
                    "rows_truncated": len(table.rows) > safe_max_rows,
                    "rows": table_rows,
                }
            )

        has_more_tables = len(doc.tables) > safe_max_tables

        if format.lower() == "json":
            res_dict: Dict[str, Any] = {
                "file_path": os.path.abspath(file_path),
                "total_tables": len(doc.tables),
                "returned_tables": len(tables_data),
                "has_more_tables": has_more_tables,
                "tables": tables_data,
            }
            if has_more_tables:
                res_dict["notice"] = (
                    f"... [TRUNCATED: Showing {len(tables_data)} of {len(doc.tables)} tables. Increase max_tables to view more] ..."
                )
            return res_dict

        md_lines = [
            f"# Tables in {os.path.basename(file_path)}",
            f"Total Tables: {len(doc.tables)} (Showing {len(tables_data)})\n",
        ]
        for t in tables_data:
            md_lines.append(f"### Table {t['table_index']} ({t['total_rows']} rows x {t['total_columns']} cols)")
            rows = t["rows"]
            if not rows:
                md_lines.append("*Empty Table*\n")
                continue

            header = rows[0]
            md_lines.append("| " + " | ".join(header) + " |")
            md_lines.append("| " + " | ".join(["---"] * len(header)) + " |")
            for r in rows[1:]:
                cells = r + [""] * (len(header) - len(r))
                md_lines.append("| " + " | ".join(cells[: len(header)]) + " |")
            if t["rows_truncated"]:
                md_lines.append(
                    f"\n*... [{t['total_rows'] - len(rows)} rows truncated. Increase max_rows to inspect more] ...*\n"
                )
            md_lines.append("\n")

        if has_more_tables:
            md_lines.append(
                f"\n*... [TRUNCATED: {len(doc.tables) - safe_max_tables} additional tables omitted to protect context window] ...*\n"
            )

        return "\n".join(md_lines)
    except Exception as e:
        if format.lower() == "json":
            return {"error": str(e)}
        return f"Error extracting tables: {str(e)}"


@mcp.tool()
def get_document_outline(file_path: str, max_headings: int = 100) -> Dict[str, Any]:
    """
    Extract the heading hierarchy and table of contents tree (H1, H2, H3, H4) with paragraph positions.

    Args:
        file_path: Absolute or relative path to the .docx or .doc file.
        max_headings: Maximum headings to extract (default: 100, max: 200).
    """
    try:
        doc = load_document(file_path)
        safe_max = min(max(1, max_headings), 200)
        headings = []

        for p_idx, p in enumerate(doc.paragraphs):
            style_name = p.style.name if p.style else ""
            text = p.text.strip()
            if not text:
                continue

            level = None
            if style_name.startswith("Heading"):
                try:
                    level = int(style_name.replace("Heading", "").strip())
                except:
                    level = 1
            elif style_name == "Title":
                level = 0
            elif style_name == "Subtitle":
                level = 1

            if level is not None:
                headings.append({"paragraph_index": p_idx, "level": level, "style": style_name, "text": text})

        has_more = len(headings) > safe_max
        displayed = headings[:safe_max]

        res: Dict[str, Any] = {
            "file_path": os.path.abspath(file_path),
            "total_headings": len(headings),
            "returned_headings": len(displayed),
            "has_more": has_more,
            "outline": displayed,
        }
        if has_more:
            res["notice"] = (
                f"... [TRUNCATED: Showing first {safe_max} headings of {len(headings)}. Increase max_headings to see more] ..."
            )
        return res
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def get_document_metadata(file_path: str) -> Dict[str, Any]:
    """
    Retrieve document properties, author, title, revision, word count, and timestamp metadata.

    Args:
        file_path: Absolute or relative path to the .docx or .doc file.
    """
    try:
        doc = load_document(file_path)
        props = doc.core_properties

        total_paras = len(doc.paragraphs)
        total_words = sum(len(p.text.split()) for p in doc.paragraphs)
        for t in doc.tables:
            seen_cells = set()
            for row in t.rows:
                for cell in row.cells:
                    if cell not in seen_cells:
                        seen_cells.add(cell)
                        total_words += len(cell.text.split())

        meta: Dict[str, Any] = {
            "file_path": os.path.abspath(file_path),
            "title": props.title or "",
            "author": props.author or "",
            "subject": props.subject or "",
            "keywords": props.keywords or "",
            "comments": props.comments or "",
            "last_modified_by": props.last_modified_by or "",
            "revision": props.revision or 1,
            "created": props.created.isoformat() if props.created else None,
            "modified": props.modified.isoformat() if props.modified else None,
            "category": props.category or "",
            "statistics": {
                "total_paragraphs": total_paras,
                "total_tables": len(doc.tables),
                "total_sections": len(doc.sections),
                "estimated_word_count": total_words,
            },
        }
        return meta
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def read_word_document(
    file_path: str,
    include_tables: bool = True,
    include_headers_footers: bool = True,
    output_format: str = "markdown",
    max_paragraphs: int = 80,
    max_chars: int = 25000,
) -> str:
    """
    Read and extract the content of a Word document formatted as Markdown or structured JSON.
    Includes token-safe truncation for large documents to prevent LLM context bloating.

    Args:
        file_path: Path to the .docx or .doc file.
        include_tables: Whether to include table contents (default: True).
        include_headers_footers: Whether to include header and footer text (default: True).
        output_format: Output format ('markdown' or 'json'). Default is 'markdown'.
        max_paragraphs: Maximum body paragraphs to extract (default: 80, max: 300).
        max_chars: Maximum characters to return before truncating (default: 25000, ~6000 tokens).
    """
    try:
        doc = load_document(file_path)
        safe_max_paras = min(max(1, max_paragraphs), 300)

        if output_format.lower() == "json":
            paragraphs_data = []
            for p_idx, p in enumerate(doc.paragraphs):
                if p.text.strip():
                    runs = []
                    for r in p.runs:
                        if r.text:
                            runs.append({"text": r.text, "formatting": extract_run_formatting(r, p.style)})
                    paragraphs_data.append(
                        {
                            "index": p_idx,
                            "style": p.style.name if p.style else "",
                            "alignment": str(p.alignment).split(".")[-1] if p.alignment else "LEFT",
                            "text": p.text.strip(),
                            "runs": runs,
                        }
                    )

            has_more = len(paragraphs_data) > safe_max_paras
            display_paras = paragraphs_data[:safe_max_paras]
            result: Dict[str, Any] = {
                "file_path": os.path.abspath(file_path),
                "total_paragraphs": len(paragraphs_data),
                "returned_paragraphs": len(display_paras),
                "has_more": has_more,
                "paragraphs": display_paras,
            }
            if has_more:
                result["notice"] = (
                    f"... [TRUNCATED: Showing first {safe_max_paras} paragraphs of {len(paragraphs_data)} to protect context window] ..."
                )
            if include_tables:
                tables_res = []
                for t in doc.tables[:10]:
                    rows = [[cell.text.strip() for cell in r.cells] for r in t.rows[:25]]
                    tables_res.append(rows)
                result["tables"] = tables_res
            return json.dumps(result, indent=2)

        md = []
        md.append(f"# Document: {os.path.basename(file_path)}\n")

        if include_headers_footers and doc.sections:
            sec = doc.sections[0]
            header_text = "\n".join(p.text.strip() for p in sec.header.paragraphs if p.text.strip())
            if header_text:
                md.append(f"> **Header**: {header_text}\n")

        paras_processed = 0
        truncated_by_paras = False

        for p in doc.paragraphs:
            if paras_processed >= safe_max_paras:
                truncated_by_paras = True
                break
            text = p.text.strip()
            if not text:
                continue

            paras_processed += 1
            style = p.style.name if p.style else "Normal"
            if style.startswith("Heading 1"):
                md.append(f"\n# {text}\n")
            elif style.startswith("Heading 2"):
                md.append(f"\n## {text}\n")
            elif style.startswith("Heading 3"):
                md.append(f"\n### {text}\n")
            elif style.startswith("Heading 4"):
                md.append(f"\n#### {text}\n")
            elif style == "Title":
                md.append(f"\n# {text}\n")
            elif style == "Subtitle":
                md.append(f"\n### *{text}*\n")
            elif "List" in style:
                md.append(f"- {text}")
            else:
                formatted_runs = []
                for r in p.runs:
                    r_text = r.text
                    if r.bold and r.italic:
                        r_text = f"***{r_text}***"
                    elif r.bold:
                        r_text = f"**{r_text}**"
                    elif r.italic:
                        r_text = f"*{r_text}*"
                    formatted_runs.append(r_text)
                md.append("".join(formatted_runs))

        if include_tables and doc.tables:
            md.append("\n## Embedded Tables\n")
            for t_idx, table in enumerate(doc.tables[:10], 1):
                md.append(f"### Table {t_idx}")
                rows = [[cell.text.strip().replace("\n", " ") for cell in r.cells] for r in table.rows]
                if rows:
                    header = rows[0]
                    md.append("| " + " | ".join(header) + " |")
                    md.append("| " + " | ".join(["---"] * len(header)) + " |")
                    for r in rows[1:25]:
                        cells = r + [""] * (len(header) - len(r))
                        md.append("| " + " | ".join(cells[: len(header)]) + " |")
                    if len(rows) > 25:
                        md.append(f"\n*... [{len(rows) - 25} table rows truncated] ...*\n")
                md.append("\n")

        if include_headers_footers and doc.sections:
            sec = doc.sections[0]
            footer_text = "\n".join(p.text.strip() for p in sec.footer.paragraphs if p.text.strip())
            if footer_text:
                md.append(f"\n> **Footer**: {footer_text}\n")

        full_output = "\n".join(md)
        if len(full_output) > max_chars:
            notice = f"\n\n... [TRUNCATED: Document exceeds {max_chars} characters. Use get_document_outline or search_word_document to inspect specific sections] ..."
            return full_output[:max_chars] + notice
        elif truncated_by_paras:
            full_output += f"\n\n... [TRUNCATED: Showing first {safe_max_paras} paragraphs of {len(doc.paragraphs)}. Increase max_paragraphs or use search_word_document to explore further] ..."
        return full_output
    except Exception as e:
        return f"Error reading Word document: {str(e)}"


@mcp.tool()
def search_word_document(
    file_path: str, search_term: str, match_case: bool = False, max_matches: int = 50
) -> Dict[str, Any]:
    """
    Search for a text term or pattern across paragraphs, headers, footers, and table cells in a Word document.

    Args:
        file_path: Path to the .docx or .doc file.
        search_term: String or pattern to search for.
        match_case: Case-sensitive search flag (default: False).
        max_matches: Maximum number of matches to return (default: 50, max: 200).
    """
    try:
        doc = load_document(file_path)
        matches = []
        flags = 0 if match_case else re.IGNORECASE
        pattern = re.compile(re.escape(search_term), flags)
        safe_max = min(max(1, max_matches), 200)

        for p_idx, p in enumerate(doc.paragraphs):
            if pattern.search(p.text):
                matches.append(
                    {
                        "location": f"Paragraph #{p_idx + 1}",
                        "style": p.style.name if p.style else "",
                        "text": p.text.strip(),
                    }
                )
                if len(matches) >= safe_max:
                    break

        if len(matches) < safe_max:
            for s_idx, section in enumerate(doc.sections, 1):
                h_text = " ".join(p.text for p in section.header.paragraphs)
                if pattern.search(h_text):
                    matches.append({"location": f"Section {s_idx} Header", "text": h_text.strip()})
                f_text = " ".join(p.text for p in section.footer.paragraphs)
                if pattern.search(f_text):
                    matches.append({"location": f"Section {s_idx} Footer", "text": f_text.strip()})
                if len(matches) >= safe_max:
                    break

        if len(matches) < safe_max:
            for t_idx, table in enumerate(doc.tables, 1):
                for r_idx, row in enumerate(table.rows):
                    for c_idx, cell in enumerate(row.cells):
                        if pattern.search(cell.text):
                            matches.append(
                                {
                                    "location": f"Table {t_idx}, Row {r_idx + 1}, Col {c_idx + 1}",
                                    "text": cell.text.strip(),
                                }
                            )
                            if len(matches) >= safe_max:
                                break
                    if len(matches) >= safe_max:
                        break
                if len(matches) >= safe_max:
                    break

        result_dict: Dict[str, Any] = {
            "file_path": os.path.abspath(file_path),
            "search_term": search_term,
            "total_matches": len(matches),
            "limit_reached": len(matches) >= safe_max,
            "matches": matches,
        }
        if len(matches) >= safe_max:
            result_dict["notice"] = (
                f"... [TRUNCATED: Showing first {safe_max} matches. Narrow your search term for more specific results] ..."
            )

        return result_dict
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def edit_paragraph(
    file_path: str, paragraph_index: int, new_text: str, output_path: Optional[str] = None, overwrite: bool = False
) -> Dict[str, Any]:
    """
    Safely edit the text of a specific paragraph (1-indexed) in a Word document, preserving its style and formatting.

    Args:
        file_path: Path to the .docx file.
        paragraph_index: 1-based index of the paragraph to edit.
        new_text: Replacement text for the paragraph.
        output_path: Optional destination path to save modified document. If omitted and overwrite=True, modifies original file.
        overwrite: Safety confirmation flag required if modifying the file directly without output_path.
    """
    try:
        doc = load_document(file_path)
        if paragraph_index < 1 or paragraph_index > len(doc.paragraphs):
            return {
                "error": f"Invalid paragraph_index {paragraph_index}. Document has {len(doc.paragraphs)} paragraphs (1-indexed)."
            }

        target_para = doc.paragraphs[paragraph_index - 1]
        old_text = target_para.text

        if target_para.runs:
            target_para.runs[0].text = new_text
            for extra_run in target_para.runs[1:]:
                extra_run.text = ""
        else:
            target_para.text = new_text

        save_dest = output_path if output_path else file_path
        if not output_path and not overwrite:
            return {"error": "Overwriting the original file requires 'overwrite=True' or specifying an 'output_path'."}

        doc.save(save_dest)
        return {
            "status": "success",
            "file_path": os.path.abspath(save_dest),
            "paragraph_index": paragraph_index,
            "old_text_preview": old_text[:100] + ("..." if len(old_text) > 100 else ""),
            "new_text_preview": new_text[:100] + ("..." if len(new_text) > 100 else ""),
        }
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def insert_table(
    file_path: str,
    headers: List[str],
    rows: List[List[str]],
    style: str = "Table Grid",
    output_path: Optional[str] = None,
    overwrite: bool = False,
) -> Dict[str, Any]:
    """
    Insert a structured table into a Word document.

    Args:
        file_path: Path to the target .docx file.
        headers: Column header names.
        rows: 2D list of row cell values.
        style: Word table style (default: 'Table Grid').
        output_path: Optional destination path. If omitted and overwrite=True, modifies original file.
        overwrite: Safety confirmation flag required if modifying the file directly without output_path.
    """
    try:
        if not headers:
            return {"error": "Headers list cannot be empty."}

        doc = load_document(file_path)
        num_cols = len(headers)

        if len(rows) > 500:
            return {"error": "Maximum 500 rows allowed per insert_table call to prevent corruption/bloat."}

        table = doc.add_table(rows=1 + len(rows), cols=num_cols)
        try:
            table.style = style
        except Exception:
            pass

        hdr_cells = table.rows[0].cells
        for col_idx, header_text in enumerate(headers):
            hdr_cells[col_idx].text = str(header_text)

        for row_idx, row_data in enumerate(rows):
            row_cells = table.rows[row_idx + 1].cells
            for col_idx in range(num_cols):
                val = row_data[col_idx] if col_idx < len(row_data) else ""
                row_cells[col_idx].text = str(val)

        save_dest = output_path if output_path else file_path
        if not output_path and not overwrite:
            return {"error": "Overwriting the original file requires 'overwrite=True' or specifying an 'output_path'."}

        doc.save(save_dest)
        return {
            "status": "success",
            "file_path": os.path.abspath(save_dest),
            "columns": num_cols,
            "rows_inserted": len(rows),
            "table_index": len(doc.tables),
        }
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def inspect_revisions_and_comments(file_path: str, max_items: int = 50) -> Dict[str, Any]:
    """
    Extract author comments, tracked insertions, and tracked deletions from a Word document (.docx).

    Args:
        file_path: Path to the .docx or .doc file.
        max_items: Maximum items to extract per category (default: 50, max: 100).
    """
    try:
        doc = load_document(file_path)
        safe_max = min(max(1, max_items), 100)
        comments = []
        insertions = []
        deletions = []

        # 1. Search for comments in docx package parts
        try:
            import xml.etree.ElementTree as ET

            for part in doc.part.package.parts:
                if "comments" in part.partname:
                    root = ET.fromstring(part.blob)
                    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                    for c_elem in root.findall(".//w:comment", ns):
                        cid = c_elem.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id", "")
                        author = c_elem.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}author", "")
                        date = c_elem.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}date", "")
                        text = "".join(c_elem.itertext()).strip()
                        comments.append({"id": cid, "author": author, "date": date, "text": text[:200]})
                        if len(comments) >= safe_max:
                            break
        except Exception:
            pass

        # 2. Search for tracked changes in body elements
        try:
            for ins in doc._element.xpath("//*[local-name()='ins']")[:safe_max]:
                author = ins.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}author", "")
                date = ins.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}date", "")
                text = "".join(ins.itertext()).strip()
                if text:
                    insertions.append({"author": author, "date": date, "text": text[:150]})
            for d in doc._element.xpath("//*[local-name()='del']")[:safe_max]:
                author = d.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}author", "")
                date = d.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}date", "")
                text = "".join(d.itertext()).strip()
                if text:
                    deletions.append({"author": author, "date": date, "text": text[:150]})
        except Exception:
            pass

        return {
            "file_path": os.path.abspath(file_path),
            "total_comments": len(comments),
            "total_insertions": len(insertions),
            "total_deletions": len(deletions),
            "comments": comments,
            "tracked_insertions": insertions,
            "tracked_deletions": deletions,
        }
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def create_document(
    file_path: str,
    title: Optional[str] = None,
    author: Optional[str] = None,
    paper_size: str = "A4",
    orientation: str = "portrait",
    margin_inches: float = 1.0,
    overwrite: bool = True,
) -> Dict[str, Any]:
    """
    Create a new empty Microsoft Word (.docx) document with custom page setup, margins, and metadata.

    Args:
        file_path: Path to the new .docx file to create.
        title: Optional document title to set in core metadata and insert as Title heading.
        author: Optional author name to set in core metadata.
        paper_size: 'A4' (default), 'Letter', 'Legal', or 'A3'.
        orientation: 'portrait' (default) or 'landscape'.
        margin_inches: Margins on all sides in inches (default: 1.0).
        overwrite: Whether to overwrite if the file already exists (default: True).
    """
    try:
        abs_path = os.path.abspath(file_path.strip("\"'"))
        if os.path.exists(abs_path) and not overwrite:
            return {"error": f"File already exists: {abs_path}. Set overwrite=True to replace it."}

        os.makedirs(os.path.dirname(abs_path), exist_ok=True)
        doc = docx.Document()
        section = doc.sections[0]

        # Page size
        p_size = paper_size.upper().strip()
        if p_size == "LETTER":
            w, h = Inches(8.5), Inches(11.0)
        elif p_size == "LEGAL":
            w, h = Inches(8.5), Inches(14.0)
        elif p_size == "A3":
            w, h = Inches(11.69), Inches(16.54)
        else:
            w, h = Inches(8.27), Inches(11.69)

        if orientation.lower() == "landscape":
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width = max(w, h)
            section.page_height = min(w, h)
        else:
            section.orientation = WD_ORIENT.PORTRAIT
            section.page_width = min(w, h)
            section.page_height = max(w, h)

        section.top_margin = Inches(margin_inches)
        section.bottom_margin = Inches(margin_inches)
        section.left_margin = Inches(margin_inches)
        section.right_margin = Inches(margin_inches)

        if author:
            doc.core_properties.author = author
        if title:
            doc.core_properties.title = title
            doc.add_heading(title, level=0)

        doc.save(abs_path)
        return {
            "status": "success",
            "file_path": abs_path,
            "title": title,
            "paper_size": paper_size,
            "orientation": orientation,
            "margin_inches": margin_inches,
            "file_size_bytes": os.path.getsize(abs_path),
        }
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def add_paragraph(
    file_path: str,
    text: str,
    style: Optional[str] = None,
    font_name: Optional[str] = None,
    font_size_pt: Optional[float] = None,
    bold: bool = False,
    italic: bool = False,
    underline: bool = False,
    color_hex: Optional[str] = None,
    alignment: Optional[str] = None,
    space_before_pt: Optional[float] = None,
    space_after_pt: Optional[float] = None,
    line_spacing: Optional[float] = None,
) -> Dict[str, Any]:
    """
    Append a styled paragraph to a Word document.

    Args:
        file_path: Path to the .docx document.
        text: Paragraph text content.
        style: Built-in style name (e.g. 'Normal', 'List Bullet', 'List Number', 'Quote').
        font_name: Font family name (e.g. 'Calibri', 'Times New Roman', 'Arial').
        font_size_pt: Font size in points (e.g. 11, 12, 14).
        bold: Whether text is bold.
        italic: Whether text is italic.
        underline: Whether text is underlined.
        color_hex: Text hex color code without '#' (e.g. '002060', 'C00000').
        alignment: 'left', 'center', 'right', or 'justify'.
        space_before_pt: Paragraph spacing before in points.
        space_after_pt: Paragraph spacing after in points.
        line_spacing: Line spacing multiple (e.g. 1.0, 1.15, 1.5, 2.0).
    """
    try:
        doc = load_document(file_path)
        p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
        run = p.add_run(text)

        if font_name:
            run.font.name = font_name
        if font_size_pt:
            run.font.size = Pt(font_size_pt)
        if bold:
            run.bold = True
        if italic:
            run.italic = True
        if underline:
            run.underline = True
        if color_hex:
            clean_hex = color_hex.lstrip("#")
            if len(clean_hex) == 6:
                r, g, b = int(clean_hex[0:2], 16), int(clean_hex[2:4], 16), int(clean_hex[4:6], 16)
                run.font.color.rgb = RGBColor(r, g, b)

        if alignment:
            align_map = {
                "center": WD_ALIGN_PARAGRAPH.CENTER,
                "right": WD_ALIGN_PARAGRAPH.RIGHT,
                "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
                "left": WD_ALIGN_PARAGRAPH.LEFT,
            }
            if alignment.lower() in align_map:
                p.alignment = align_map[alignment.lower()]

        pf = p.paragraph_format
        if space_before_pt is not None:
            pf.space_before = Pt(space_before_pt)
        if space_after_pt is not None:
            pf.space_after = Pt(space_after_pt)
        if line_spacing is not None:
            pf.line_spacing = line_spacing

        clean_path = os.path.abspath(file_path.strip("\"'"))
        doc.save(clean_path)
        return {
            "status": "success",
            "file_path": clean_path,
            "paragraph_index": len(doc.paragraphs) - 1,
            "text_length": len(text),
        }
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def add_heading(file_path: str, text: str, level: int = 1) -> Dict[str, Any]:
    """
    Append a heading (level 1-9) or Title (level 0) to a Word document.

    Args:
        file_path: Path to the .docx document.
        text: Heading text content.
        level: Heading level from 0 (Title) to 9 (default: 1).
    """
    try:
        doc = load_document(file_path)
        safe_level = max(0, min(9, int(level)))
        doc.add_heading(text, level=safe_level)
        clean_path = os.path.abspath(file_path.strip("\"'"))
        doc.save(clean_path)
        return {"status": "success", "file_path": clean_path, "heading": text, "level": safe_level}
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def fill_template(template_path: str, output_path: str, replacements: Dict[str, str]) -> Dict[str, Any]:
    """
    Populate a Word document template by replacing placeholder tags (e.g. {{client_name}}, {{date}}, {{total_amount}})
    across body paragraphs, tables, headers, and footers while preserving styles.

    Args:
        template_path: Path to the source .docx template file.
        output_path: Destination path for the populated .docx file.
        replacements: Key-value dictionary of placeholders to replacement strings (e.g. {"{{client_name}}": "Acme Inc"}).
    """
    try:
        doc = load_document(template_path)
        total_replaced = 0

        def replace_in_p(p):
            nonlocal total_replaced
            for old_token, new_val in replacements.items():
                if old_token in p.text:
                    replaced_in_runs = False
                    for run in p.runs:
                        if old_token in run.text:
                            run.text = run.text.replace(old_token, str(new_val))
                            replaced_in_runs = True
                            total_replaced += 1
                    if not replaced_in_runs and old_token in p.text:
                        p.text = p.text.replace(old_token, str(new_val))
                        total_replaced += 1

        for p in doc.paragraphs:
            replace_in_p(p)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        replace_in_p(p)

        for sec in doc.sections:
            for h in (sec.header, sec.first_page_header, sec.even_page_header):
                if h:
                    for p in h.paragraphs:
                        replace_in_p(p)
            for f in (sec.footer, sec.first_page_footer, sec.even_page_footer):
                if f:
                    for p in f.paragraphs:
                        replace_in_p(p)

        dest = os.path.abspath(output_path.strip("\"'"))
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        doc.save(dest)
        return {
            "status": "success",
            "template_path": os.path.abspath(template_path),
            "output_path": dest,
            "placeholders_count": len(replacements),
            "total_replacements_applied": total_replaced,
        }
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def replace_text(
    file_path: str, search_text: str, replace_text: str, match_case: bool = False, output_path: Optional[str] = None
) -> Dict[str, Any]:
    """
    Find and replace text across the entire Word document (paragraphs and tables).

    Args:
        file_path: Path to the .docx document.
        search_text: Substring to find.
        replace_text: Replacement substring.
        match_case: Whether to match case sensitive (default: False).
        output_path: Optional output path (overwrites file_path if omitted).
    """
    try:
        doc = load_document(file_path)
        replacements_count = 0
        flags = 0 if match_case else re.IGNORECASE
        pattern = re.compile(re.escape(search_text), flags)

        def do_replace(p):
            nonlocal replacements_count
            if pattern.search(p.text):
                count_in_p = len(pattern.findall(p.text))
                p.text = pattern.sub(replace_text, p.text)
                replacements_count += count_in_p

        for p in doc.paragraphs:
            do_replace(p)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        do_replace(p)

        dest = os.path.abspath((output_path or file_path).strip("\"'"))
        doc.save(dest)
        return {
            "status": "success",
            "file_path": dest,
            "search_text": search_text,
            "occurrences_replaced": replacements_count,
        }
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def insert_image(
    file_path: str,
    image_path: str,
    width_inches: Optional[float] = None,
    height_inches: Optional[float] = None,
    caption: Optional[str] = None,
) -> Dict[str, Any]:
    """
    Insert an image into a Word document with optional scaling and caption.

    Args:
        file_path: Path to the target .docx document.
        image_path: Path to the image file (.png, .jpg, .jpeg, etc.).
        width_inches: Image width in inches (preserves aspect ratio if height omitted).
        height_inches: Image height in inches (preserves aspect ratio if width omitted).
        caption: Optional caption text placed below the image.
    """
    try:
        clean_img = os.path.abspath(image_path.strip("\"'"))
        if not os.path.exists(clean_img):
            return {"error": f"Image file not found: {clean_img}"}

        doc = load_document(file_path)
        w = Inches(width_inches) if width_inches else None
        h = Inches(height_inches) if height_inches else None

        doc.add_picture(clean_img, width=w, height=h)
        if caption:
            p = doc.add_paragraph(caption, style="Caption")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        dest = os.path.abspath(file_path.strip("\"'"))
        doc.save(dest)
        return {
            "status": "success",
            "file_path": dest,
            "image_path": clean_img,
            "width_inches": width_inches,
            "height_inches": height_inches,
            "caption": caption,
        }
    except Exception as e:
        return {"error": str(e)}


@mcp.tool()
def export_to_pdf(file_path: str, output_pdf_path: Optional[str] = None) -> Dict[str, Any]:
    """
    Export a Microsoft Word document (.docx) to a high-fidelity PDF file using Windows Word COM automation.

    Args:
        file_path: Path to the source .docx file.
        output_pdf_path: Optional destination path for the .pdf file (defaults to same name with .pdf extension).
    """
    try:
        abs_docx = os.path.abspath(file_path.strip("\"'"))
        if not os.path.exists(abs_docx):
            return {"error": f"Source document not found: {abs_docx}"}

        if output_pdf_path:
            abs_pdf = os.path.abspath(output_pdf_path.strip("\"'"))
        else:
            abs_pdf = os.path.splitext(abs_docx)[0] + ".pdf"

        os.makedirs(os.path.dirname(abs_pdf), exist_ok=True)

        try:
            import win32com.client

            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            word.DisplayAlerts = False
            try:
                wdoc = word.Documents.Open(abs_docx)
                wdoc.SaveAs2(abs_pdf, FileFormat=17)  # 17 = wdFormatPDF
                wdoc.Close(False)
            finally:
                word.Quit()
        except Exception as com_err:
            try:
                docx2pdf = __import__("docx2pdf")
                docx2pdf.convert(abs_docx, abs_pdf)
            except Exception:
                return {"error": f"Failed to export PDF via Word COM: {com_err}. Ensure Microsoft Word is installed."}

        return {
            "status": "success",
            "source_docx": abs_docx,
            "output_pdf": abs_pdf,
            "pdf_size_bytes": os.path.getsize(abs_pdf),
        }
    except Exception as e:
        return {"error": str(e)}


if __name__ == "__main__":
    mcp.run()
