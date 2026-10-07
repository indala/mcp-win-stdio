import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
from pathlib import Path

# Test 1: create_document
doc = docx.Document()
section = doc.sections[0]
section.top_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)

# Test 2: add_heading
doc.add_heading("Quarterly Performance Report", level=1)

# Test 3: add_paragraph
p = doc.add_paragraph()
run = p.add_run("This report summarizes the operational KPIs for Q3 2026. Hello {{company_name}}!")
run.font.name = "Calibri"
run.font.size = Pt(11)
run.bold = False

# Test 4: table
tbl = doc.add_table(rows=1, cols=3)
hdr = tbl.rows[0].cells
hdr[0].text = "Department"
hdr[1].text = "Status"
hdr[2].text = "Revenue"

row = tbl.add_row().cells
row[0].text = "Props Management"
row[1].text = "Active"
row[2].text = "$1,240,000"

out_path = "scratch/test_word_doc.docx"
doc.save(out_path)
print(f"[PASS] Document created successfully at {out_path}, size: {os.path.getsize(out_path)} bytes")

# Test 5: fill_template
def replace_in_p(paragraph, old, new):
    if old in paragraph.text:
        # replace across runs or full text
        paragraph.text = paragraph.text.replace(old, new)

doc2 = docx.Document(out_path)
for p in doc2.paragraphs:
    replace_in_p(p, "{{company_name}}", "Acme Production Studios")

out_filled = "scratch/test_word_filled.docx"
doc2.save(out_filled)
print(f"[PASS] Template filled successfully at {out_filled}")
doc3 = docx.Document(out_filled)
text = " ".join([p.text for p in doc3.paragraphs])
assert "Acme Production Studios" in text
print("[PASS] Template token replacement verified!")
