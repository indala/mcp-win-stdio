#!/usr/bin/env python3
"""
Word Document MCP Server for Windows (Advanced Layout & Typography Edition)
Extracts comprehensive content, headers, footers, advanced page geometry, exact margins across all units,
multi-column layouts (IEEE/Journal), paragraph spacing/indentation, typography, tables, outline, images, and metadata.
"""

import os
import sys
import json
import re
from typing import Optional, List, Dict, Any
from mcp.server.fastmcp import FastMCP
import docx
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Inches, Pt, Cm
from docx.oxml.ns import nsmap

# Register namespaces
nsmap['wp'] = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
nsmap['a'] = 'http://schemas.openxmlformats.org/drawingml/2006/main'
nsmap['pic'] = 'http://schemas.openxmlformats.org/drawingml/2006/picture'
nsmap['r'] = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
nsmap['v'] = 'urn:schemas-microsoft-com:vml'

mcp = FastMCP("word-mcp")

# Unit Conversion Constants
EMU_PER_INCH = 914400.0
EMU_PER_CM = 360000.0
EMU_PER_PT = 12700.0
TWIPS_PER_INCH = 1440.0
TWIPS_PER_CM = 567.0
TWIPS_PER_PT = 20.0

def identify_paper_size(width_in: float, height_in: float) -> str:
    w, h = sorted([round(width_in, 2), round(height_in, 2)])
    if abs(w - 8.27) < 0.05 and abs(h - 11.69) < 0.05:
        return "A4 (8.27 x 11.69 in / 210 x 297 mm)"
    elif abs(w - 8.5) < 0.05 and abs(h - 11.0) < 0.05:
        return "Letter (8.5 x 11 in / 216 x 279 mm)"
    elif abs(w - 8.5) < 0.05 and abs(h - 14.0) < 0.05:
        return "Legal (8.5 x 14 in / 216 x 356 mm)"
    elif abs(w - 5.83) < 0.05 and abs(h - 8.27) < 0.05:
        return "A5 (5.83 x 8.27 in / 148 x 210 mm)"
    elif abs(w - 11.69) < 0.05 and abs(h - 16.54) < 0.05:
        return "A3 (11.69 x 16.54 in / 297 x 420 mm)"
    elif abs(w - 7.25) < 0.05 and abs(h - 10.5) < 0.05:
        return "Executive (7.25 x 10.5 in / 184 x 267 mm)"
    elif abs(w - 7.17) < 0.05 and abs(h - 10.12) < 0.05:
        return "B5 JIS (7.17 x 10.12 in / 182 x 257 mm)"
    return f"Custom ({width_in:.2f} x {height_in:.2f} in)"

def load_document(file_path: str) -> docx.Document:
    clean_path = os.path.abspath(file_path.strip('"\''))
    if not os.path.exists(clean_path):
        raise FileNotFoundError(f"File not found: {clean_path}")
    
    if clean_path.lower().endswith(".doc") and not clean_path.lower().endswith(".docx"):
        try:
            import win32com.client
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            doc = word.Documents.Open(clean_path)
            temp_docx = clean_path + "x"
            doc.SaveAs2(temp_docx, FileFormat=16)
            doc.Close()
            word.Quit()
            return docx.Document(temp_docx)
        except Exception as e:
            raise RuntimeError(f"Could not convert legacy .doc file using MS Word: {e}")

    return docx.Document(clean_path)

def extract_run_formatting(run, paragraph_style=None) -> Dict[str, Any]:
    """Extract font name, size (pt), color (hex/theme), highlight, and bold/italic."""
    font_name = run.font.name
    if not font_name and paragraph_style and hasattr(paragraph_style, 'font'):
        font_name = paragraph_style.font.name

    font_size = None
    if run.font.size:
        font_size = run.font.size.pt
    elif paragraph_style and hasattr(paragraph_style, 'font') and paragraph_style.font.size:
        font_size = paragraph_style.font.size.pt

    color_hex = None
    if run.font.color and run.font.color.rgb:
        color_hex = f"#{run.font.color.rgb}"

    highlight = str(run.font.highlight_color) if run.font.highlight_color else None

    rPr = run._r.xpath('.//*[local-name()="rPr"]')
    if rPr:
        if not color_hex:
            color_el = rPr[0].xpath('.//*[local-name()="color"]')
            if color_el:
                val = color_el[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') or color_el[0].get('val')
                if val and val != 'auto':
                    color_hex = f"#{val}"
                theme = color_el[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}themeColor') or color_el[0].get('themeColor')
                if theme and not color_hex:
                    color_hex = f"theme:{theme}"
        if not font_size:
            sz_el = rPr[0].xpath('.//*[local-name()="sz"]')
            if sz_el:
                val = sz_el[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') or sz_el[0].get('val')
                if val:
                    try:
                        font_size = int(val) / 2.0
                    except:
                        pass
        if not font_name:
            rFonts_el = rPr[0].xpath('.//*[local-name()="rFonts"]')
            if rFonts_el:
                font_name = rFonts_el[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii') or rFonts_el[0].get('ascii')

    return {
        "font_family": font_name or "Default / Body Font",
        "font_size_pt": font_size if font_size is not None else "Default / 11pt",
        "font_color": color_hex or "#000000 (Automatic)",
        "highlight": highlight,
        "bold": bool(run.bold),
        "italic": bool(run.italic),
        "underline": bool(run.underline)
    }

def extract_drawing_info(element, location: str, part=None) -> List[Dict[str, Any]]:
    images = []
    if element is None:
        return images

    drawings = element.xpath('.//*[local-name()="drawing"]')

    for d in drawings:
        anchors = d.xpath('.//*[local-name()="anchor"]')
        inlines = d.xpath('.//*[local-name()="inline"]')

        for item in (anchors + inlines):
            is_anchor = item in anchors
            placement_type = "FLOATING_ANCHOR" if is_anchor else "INLINE"

            ext = item.xpath('.//*[local-name()="extent"]')
            cx = int(ext[0].get('cx', 0)) if ext else 0
            cy = int(ext[0].get('cy', 0)) if ext else 0

            doc_pr = item.xpath('.//*[local-name()="docPr"]')
            name = doc_pr[0].get('name', '') if doc_pr else ''
            descr = doc_pr[0].get('descr', '') if doc_pr else ''
            title = doc_pr[0].get('title', '') if doc_pr else ''

            blips = item.xpath('.//*[local-name()="blip"]')
            r_id = ''
            if blips:
                for attr_name, attr_val in blips[0].attrib.items():
                    if attr_name.endswith('embed'):
                        r_id = attr_val
                        break

            target_file = ''
            content_type = ''
            file_size_bytes = None

            if r_id and part and hasattr(part, 'rels'):
                try:
                    rel = part.rels.get(r_id)
                    if rel:
                        target_file = os.path.basename(rel.target_ref)
                        if hasattr(part, 'related_parts'):
                            related_part = part.related_parts.get(r_id)
                            if related_part:
                                content_type = getattr(related_part, 'content_type', '')
                                blob = getattr(related_part, 'blob', None)
                                if blob:
                                    file_size_bytes = len(blob)
                except Exception:
                    pass

            h_align = None
            h_rel_from = None
            h_offset_in = None
            pos_h = item.xpath('.//*[local-name()="positionH"]')
            if pos_h:
                h_rel_from = pos_h[0].get('relativeFrom', '')
                align_el = pos_h[0].xpath('.//*[local-name()="align"]')
                if align_el and align_el[0].text:
                    h_align = align_el[0].text.strip()
                pos_offset = pos_h[0].xpath('.//*[local-name()="posOffset"]')
                if pos_offset and pos_offset[0].text:
                    try:
                        h_offset_in = round(int(pos_offset[0].text) / EMU_PER_INCH, 2)
                    except:
                        pass

            v_align = None
            v_rel_from = None
            v_offset_in = None
            pos_v = item.xpath('.//*[local-name()="positionV"]')
            if pos_v:
                v_rel_from = pos_v[0].get('relativeFrom', '')
                align_el = pos_v[0].xpath('.//*[local-name()="align"]')
                if align_el and align_el[0].text:
                    v_align = align_el[0].text.strip()
                pos_offset = pos_v[0].xpath('.//*[local-name()="posOffset"]')
                if pos_offset and pos_offset[0].text:
                    try:
                        v_offset_in = round(int(pos_offset[0].text) / EMU_PER_INCH, 2)
                    except:
                        pass

            images.append({
                "location": location,
                "placement_type": placement_type,
                "image_name": name or "Unnamed Image",
                "alt_text": descr or title or None,
                "dimensions": {
                    "width_inches": round(cx / EMU_PER_INCH, 2),
                    "height_inches": round(cy / EMU_PER_INCH, 2),
                    "width_cm": round(cx / EMU_PER_CM, 2),
                    "height_cm": round(cy / EMU_PER_CM, 2),
                    "width_pt": round(cx / EMU_PER_PT, 1),
                    "height_pt": round(cy / EMU_PER_PT, 1)
                },
                "positioning": {
                    "horizontal_alignment": h_align or ("offset" if h_offset_in is not None else "inline"),
                    "horizontal_relative_from": h_rel_from or None,
                    "horizontal_offset_inches": h_offset_in,
                    "vertical_alignment": v_align or ("offset" if v_offset_in is not None else "inline"),
                    "vertical_relative_from": v_rel_from or None,
                    "vertical_offset_inches": v_offset_in
                },
                "embedded_file": {
                    "target_filename": target_file or None,
                    "content_type": content_type or None,
                    "file_size_bytes": file_size_bytes
                }
            })

    return images


@mcp.tool()
def get_document_layout(file_path: str) -> str:
    """
    Extract comprehensive page layout geometry, exact margins across multiple units (in, cm, mm, pt, twips),
    multi-column layout (IEEE/Journal), printable area, section break types, and header/footer distances.
    
    Args:
        file_path: Absolute or relative path to the .docx or .doc file.
    """
    try:
        doc = load_document(file_path)
        sections_info = []

        for idx, section in enumerate(doc.sections, 1):
            sectPr = section._sectPr

            # Dimensions
            w_in = section.page_width.inches if section.page_width else 8.5
            h_in = section.page_height.inches if section.page_height else 11.0
            orientation_str = "PORTRAIT" if section.orientation == WD_ORIENT.PORTRAIT else "LANDSCAPE"
            paper_size = identify_paper_size(w_in, h_in)

            # Margins
            top_in = section.top_margin.inches if section.top_margin else 1.0
            bot_in = section.bottom_margin.inches if section.bottom_margin else 1.0
            left_in = section.left_margin.inches if section.left_margin else 1.0
            right_in = section.right_margin.inches if section.right_margin else 1.0
            gutter_in = section.gutter.inches if section.gutter else 0.0
            header_dist_in = section.header_distance.inches if section.header_distance else 0.5
            footer_dist_in = section.footer_distance.inches if section.footer_distance else 0.5

            # Printable Area
            printable_w_in = round(w_in - left_in - right_in - gutter_in, 2)
            printable_h_in = round(h_in - top_in - bot_in, 2)

            # Section Type & Flow
            type_el = sectPr.xpath('.//*[local-name()="type"]')
            sec_type_raw = type_el[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') or type_el[0].get('val') if type_el else 'nextPage'
            sec_type_map = {
                'nextPage': 'NEW_PAGE (Starts on Next Page)',
                'continuous': 'CONTINUOUS (Continuous Section Break)',
                'evenPage': 'EVEN_PAGE (Starts on Next Even Page)',
                'oddPage': 'ODD_PAGE (Starts on Next Odd Page)',
                'nextColumn': 'NEXT_COLUMN (Starts on Next Column)'
            }
            sec_type_readable = sec_type_map.get(sec_type_raw, sec_type_raw)

            # Multi-Column Layout (IEEE / Journal 2-Column layout)
            cols_el = sectPr.xpath('.//*[local-name()="cols"]')
            num_cols = 1
            col_space_in = None
            equal_width = True
            line_sep = False
            col_width_in = printable_w_in

            if cols_el:
                num_attr = cols_el[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num') or cols_el[0].get('num')
                if num_attr:
                    try:
                        num_cols = int(num_attr)
                    except:
                        num_cols = 1
                space_attr = cols_el[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}space') or cols_el[0].get('space')
                if space_attr:
                    try:
                        col_space_in = round(int(space_attr) / TWIPS_PER_INCH, 2)
                    except:
                        pass
                sep_attr = cols_el[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sep') or cols_el[0].get('sep')
                line_sep = bool(sep_attr and sep_attr in ['1', 'true', 'on'])
                eq_attr = cols_el[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}equalWidth') or cols_el[0].get('equalWidth')
                equal_width = not (eq_attr in ['0', 'false', 'off'])

                if num_cols > 1:
                    spacing_total = (col_space_in or 0.25) * (num_cols - 1)
                    col_width_in = round((printable_w_in - spacing_total) / num_cols, 2)

            # Page Numbering
            pgNum_el = sectPr.xpath('.//*[local-name()="pgNumType"]')
            pg_start = None
            pg_format = 'decimal (1, 2, 3...)'
            if pgNum_el:
                start_attr = pgNum_el[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}start') or pgNum_el[0].get('start')
                if start_attr:
                    try:
                        pg_start = int(start_attr)
                    except:
                        pass
                fmt_attr = pgNum_el[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fmt') or pgNum_el[0].get('fmt')
                if fmt_attr:
                    pg_format = fmt_attr

            sections_info.append({
                "section_number": idx,
                "section_break_type": sec_type_readable,
                "orientation": orientation_str,
                "paper_size": paper_size,
                "page_dimensions": {
                    "inches": {"width": round(w_in, 2), "height": round(h_in, 2)},
                    "centimeters": {"width": round(w_in * 2.54, 2), "height": round(h_in * 2.54, 2)},
                    "millimeters": {"width": round(w_in * 25.4, 1), "height": round(h_in * 25.4, 1)},
                    "points": {"width": round(w_in * 72.0, 1), "height": round(h_in * 72.0, 1)}
                },
                "printable_area": {
                    "inches": {"width": printable_w_in, "height": printable_h_in},
                    "centimeters": {"width": round(printable_w_in * 2.54, 2), "height": round(printable_h_in * 2.54, 2)},
                    "points": {"width": round(printable_w_in * 72.0, 1), "height": round(printable_h_in * 72.0, 1)}
                },
                "margins": {
                    "inches": {
                        "top": round(top_in, 2),
                        "bottom": round(bot_in, 2),
                        "left": round(left_in, 2),
                        "right": round(right_in, 2),
                        "gutter": round(gutter_in, 2)
                    },
                    "centimeters": {
                        "top": round(top_in * 2.54, 2),
                        "bottom": round(bot_in * 2.54, 2),
                        "left": round(left_in * 2.54, 2),
                        "right": round(right_in * 2.54, 2),
                        "gutter": round(gutter_in * 2.54, 2)
                    },
                    "millimeters": {
                        "top": round(top_in * 25.4, 1),
                        "bottom": round(bot_in * 25.4, 1),
                        "left": round(left_in * 25.4, 1),
                        "right": round(right_in * 25.4, 1),
                        "gutter": round(gutter_in * 25.4, 1)
                    },
                    "points": {
                        "top": round(top_in * 72.0, 1),
                        "bottom": round(bot_in * 72.0, 1),
                        "left": round(left_in * 72.0, 1),
                        "right": round(right_in * 72.0, 1),
                        "gutter": round(gutter_in * 72.0, 1)
                    }
                },
                "multi_column_layout": {
                    "is_multi_column": num_cols > 1,
                    "column_count": num_cols,
                    "column_width_inches": col_width_in,
                    "column_width_cm": round(col_width_in * 2.54, 2),
                    "column_spacing_inches": col_space_in or (0.25 if num_cols > 1 else 0.0),
                    "column_spacing_cm": round((col_space_in or (0.25 if num_cols > 1 else 0.0)) * 2.54, 2),
                    "equal_width_columns": equal_width,
                    "line_separator_between_columns": line_sep
                },
                "header_footer_geometry": {
                    "header_distance_from_top_edge_inches": round(header_dist_in, 2),
                    "header_distance_from_top_edge_cm": round(header_dist_in * 2.54, 2),
                    "footer_distance_from_bottom_edge_inches": round(footer_dist_in, 2),
                    "footer_distance_from_bottom_edge_cm": round(footer_dist_in * 2.54, 2),
                    "different_first_page": section.different_first_page_header_footer,
                    "different_odd_and_even_pages": getattr(doc.settings, 'odd_and_even_pages_header_footer', False) if hasattr(doc, 'settings') else False
                },
                "page_numbering": {
                    "starts_at": pg_start or "Continue from previous section",
                    "number_format": pg_format
                }
            })

        return json.dumps({
            "file_path": os.path.abspath(file_path),
            "total_sections": len(doc.sections),
            "sections": sections_info
        }, indent=2)
    except Exception as e:
        return json.dumps({"error": str(e)}, indent=2)


@mcp.tool()
def get_paragraph_spacing_and_indentation(file_path: str, max_paragraphs: int = 40) -> str:
    """
    Extract detailed paragraph spacing (line spacing, space before/after in pt), indentation (first-line indent, left/right margins in inches), and alignment.
    
    Args:
        file_path: Absolute or relative path to the .docx or .doc file.
        max_paragraphs: Number of paragraphs to analyze (default: 40).
    """
    try:
        doc = load_document(file_path)
        paragraphs_info = []

        for idx, p in enumerate(doc.paragraphs[:max_paragraphs]):
            text = p.text.strip()
            if not text:
                continue

            pf = p.paragraph_format
            align = str(p.alignment).split(".")[-1] if p.alignment else "LEFT"

            # Line Spacing
            line_spacing = pf.line_spacing
            line_spacing_rule = str(pf.line_spacing_rule).split(".")[-1] if pf.line_spacing_rule else "MULTIPLE"
            if line_spacing is None:
                line_spacing_desc = "Single (1.0 / Default)"
            elif isinstance(line_spacing, float):
                line_spacing_desc = f"{line_spacing:.2f}x Line Spacing"
            elif hasattr(line_spacing, 'pt'):
                line_spacing_desc = f"Exact {line_spacing.pt:.1f} pt"
            else:
                line_spacing_desc = str(line_spacing)

            # Space Before & After
            space_before_pt = pf.space_before.pt if pf.space_before else 0.0
            space_after_pt = pf.space_after.pt if pf.space_after else 0.0

            # Indentation
            first_line_in = round(pf.first_line_indent.inches, 2) if pf.first_line_indent else 0.0
            left_indent_in = round(pf.left_indent.inches, 2) if pf.left_indent else 0.0
            right_indent_in = round(pf.right_indent.inches, 2) if pf.right_indent else 0.0

            paragraphs_info.append({
                "paragraph_index": idx + 1,
                "style": p.style.name if p.style else "Normal",
                "alignment": align,
                "text_preview": text[:70] + ("..." if len(text) > 70 else ""),
                "spacing": {
                    "line_spacing": line_spacing_desc,
                    "line_spacing_rule": line_spacing_rule,
                    "space_before_pt": space_before_pt,
                    "space_after_pt": space_after_pt
                },
                "indentation": {
                    "first_line_indent_inches": first_line_in,
                    "first_line_indent_cm": round(first_line_in * 2.54, 2),
                    "left_indent_inches": left_indent_in,
                    "left_indent_cm": round(left_indent_in * 2.54, 2),
                    "right_indent_inches": right_indent_in,
                    "is_hanging_indent": first_line_in < 0
                },
                "pagination_controls": {
                    "keep_with_next": bool(pf.keep_with_next),
                    "page_break_before": bool(pf.page_break_before),
                    "widow_control": bool(pf.widow_control)
                }
            })

        return json.dumps({
            "file_path": os.path.abspath(file_path),
            "analyzed_paragraphs_count": len(paragraphs_info),
            "paragraphs": paragraphs_info
        }, indent=2)
    except Exception as e:
        return json.dumps({"error": str(e)}, indent=2)


@mcp.tool()
def get_document_typography(file_path: str) -> str:
    """
    Extract all typography details (font families, font sizes, colors, headings, styles, and alignments) across the entire Word document.
    
    Args:
        file_path: Path to the .docx or .doc file.
    """
    try:
        doc = load_document(file_path)
        
        distinct_fonts = set()
        distinct_sizes = set()
        distinct_colors = set()
        styles_summary = []

        for p_idx, p in enumerate(doc.paragraphs):
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name if p.style else "Normal"
            align = str(p.alignment).split(".")[-1] if p.alignment else "LEFT"

            runs_data = []
            for r in p.runs:
                if r.text:
                    fmt = extract_run_formatting(r, p.style)
                    distinct_fonts.add(fmt["font_family"])
                    distinct_sizes.add(str(fmt["font_size_pt"]))
                    distinct_colors.add(fmt["font_color"])
                    runs_data.append({
                        "text": r.text,
                        "formatting": fmt
                    })

            if runs_data:
                styles_summary.append({
                    "paragraph_index": p_idx + 1,
                    "style": style_name,
                    "alignment": align,
                    "text_preview": text[:80] + ("..." if len(text) > 80 else ""),
                    "runs": runs_data
                })

        headers_typography = []
        for s_idx, sec in enumerate(doc.sections, 1):
            if sec.header:
                for p_idx, p in enumerate(sec.header.paragraphs):
                    if p.text.strip():
                        runs = [extract_run_formatting(r, p.style) for r in p.runs if r.text.strip()]
                        for r_fmt in runs:
                            distinct_fonts.add(r_fmt["font_family"])
                            distinct_sizes.add(str(r_fmt["font_size_pt"]))
                            distinct_colors.add(r_fmt["font_color"])
                        headers_typography.append({
                            "location": f"Section {s_idx} Header P#{p_idx+1}",
                            "text": p.text.strip(),
                            "formatting": runs
                        })

        return json.dumps({
            "file_path": os.path.abspath(file_path),
            "summary": {
                "font_families_used": sorted(list(distinct_fonts)),
                "font_sizes_used_pt": sorted(list(distinct_sizes)),
                "colors_used": sorted(list(distinct_colors))
            },
            "header_typography": headers_typography,
            "paragraphs_typography": styles_summary[:30]
        }, indent=2)
    except Exception as e:
        return json.dumps({"error": str(e)}, indent=2)


@mcp.tool()
def get_document_images(file_path: str) -> str:
    """
    Extract all images and graphics from headers, footers, body paragraphs, and tables along with exact placement, alignment, dimensions, and coordinates.
    
    Args:
        file_path: Absolute or relative path to the .docx or .doc file.
    """
    try:
        doc = load_document(file_path)
        all_images = []

        for s_idx, sec in enumerate(doc.sections, 1):
            if sec.header:
                all_images.extend(extract_drawing_info(sec.header._element, f"Section {s_idx} Primary Header", sec.header.part))
            if sec.footer:
                all_images.extend(extract_drawing_info(sec.footer._element, f"Section {s_idx} Primary Footer", sec.footer.part))

            if sec.different_first_page_header_footer:
                if sec.first_page_header:
                    all_images.extend(extract_drawing_info(sec.first_page_header._element, f"Section {s_idx} First Page Header", sec.first_page_header.part))
                if sec.first_page_footer:
                    all_images.extend(extract_drawing_info(sec.first_page_footer._element, f"Section {s_idx} First Page Footer", sec.first_page_footer.part))

            if getattr(sec, 'even_page_header', None):
                all_images.extend(extract_drawing_info(sec.even_page_header._element, f"Section {s_idx} Even Page Header", sec.even_page_header.part))
            if getattr(sec, 'even_page_footer', None):
                all_images.extend(extract_drawing_info(sec.even_page_footer._element, f"Section {s_idx} Even Page Footer", sec.even_page_footer.part))

        for p_idx, p in enumerate(doc.paragraphs):
            p_images = extract_drawing_info(p._element, f"Body Paragraph #{p_idx + 1}", doc.part)
            all_images.extend(p_images)

        for t_idx, table in enumerate(doc.tables, 1):
            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    cell_images = extract_drawing_info(cell._tc, f"Table {t_idx}, Cell ({r_idx + 1}, {c_idx + 1})", doc.part)
                    all_images.extend(cell_images)

        return json.dumps({
            "file_path": os.path.abspath(file_path),
            "total_images_found": len(all_images),
            "header_images_count": sum(1 for img in all_images if "Header" in img["location"]),
            "footer_images_count": sum(1 for img in all_images if "Footer" in img["location"]),
            "body_images_count": sum(1 for img in all_images if "Body" in img["location"] or "Table" in img["location"]),
            "images": all_images
        }, indent=2)
    except Exception as e:
        return json.dumps({"error": str(e)}, indent=2)


@mcp.tool()
def get_headers_and_footers(file_path: str) -> str:
    """
    Extract all headers and footers from each section of a Word document, including text, typography (fonts, sizes, colors), tables, and images.
    
    Args:
        file_path: Absolute or relative path to the .docx or .doc file.
    """
    try:
        doc = load_document(file_path)
        sections_hf = []

        for idx, section in enumerate(doc.sections, 1):
            def extract_hf_data(hf_obj, loc_name):
                if not hf_obj:
                    return None
                
                paras_info = []
                for p in hf_obj.paragraphs:
                    if p.text.strip():
                        runs = [{"text": r.text, "formatting": extract_run_formatting(r, p.style)} for r in p.runs if r.text.strip()]
                        paras_info.append({
                            "text": p.text.strip(),
                            "alignment": str(p.alignment).split(".")[-1] if p.alignment else "LEFT",
                            "runs": runs
                        })

                images = extract_drawing_info(hf_obj._element, loc_name, hf_obj.part)
                return {
                    "paragraphs": paras_info,
                    "image_count": len(images),
                    "images": images
                }

            primary_h = extract_hf_data(section.header, f"Section {idx} Primary Header")
            primary_f = extract_hf_data(section.footer, f"Section {idx} Primary Footer")

            first_h = extract_hf_data(section.first_page_header, f"Section {idx} First Page Header") if section.different_first_page_header_footer else None
            first_f = extract_hf_data(section.first_page_footer, f"Section {idx} First Page Footer") if section.different_first_page_header_footer else None

            even_h = extract_hf_data(getattr(section, 'even_page_header', None), f"Section {idx} Even Page Header") if getattr(section, 'even_page_header', None) else None
            even_f = extract_hf_data(getattr(section, 'even_page_footer', None), f"Section {idx} Even Page Footer") if getattr(section, 'even_page_footer', None) else None

            sections_hf.append({
                "section_number": idx,
                "different_first_page": section.different_first_page_header_footer,
                "primary_header": primary_h,
                "primary_footer": primary_f,
                "first_page_header": first_h,
                "first_page_footer": first_f,
                "even_page_header": even_h,
                "even_page_footer": even_f
            })

        return json.dumps({
            "file_path": os.path.abspath(file_path),
            "sections": sections_hf
        }, indent=2)
    except Exception as e:
        return json.dumps({"error": str(e)}, indent=2)


@mcp.tool()
def get_document_tables(file_path: str, format: str = "markdown") -> str:
    """
    Extract all tables from a Word document as structured JSON or readable Markdown.
    
    Args:
        file_path: Absolute or relative path to the .docx or .doc file.
        format: Output format ('markdown' or 'json'). Default is 'markdown'.
    """
    try:
        doc = load_document(file_path)
        tables_data = []

        for t_idx, table in enumerate(doc.tables, 1):
            table_rows = []
            for row in table.rows:
                row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                table_rows.append(row_cells)

            tables_data.append({
                "table_index": t_idx,
                "total_rows": len(table.rows),
                "total_columns": len(table.columns),
                "rows": table_rows
            })

        if format.lower() == "json":
            return json.dumps({
                "file_path": os.path.abspath(file_path),
                "total_tables": len(tables_data),
                "tables": tables_data
            }, indent=2)

        md_lines = [f"# Tables in {os.path.basename(file_path)}", f"Total Tables: {len(tables_data)}\n"]
        for t in tables_data:
            md_lines.append(f"### Table {t['table_index']} ({t['total_rows']} rows x {t['total_columns']} cols)")
            rows = t["rows"]
            if not rows:
                md_lines.append("*Empty Table*\n")
                continue

            header = rows[0]
            md_lines.append("| " + " | ".join(header) + " |")
            md_lines.append("| " + " | ".join(["---"] * len(header)) + " |")
            for r in rows[1:]:
                cells = r + [""] * (len(header) - len(r))
                md_lines.append("| " + " | ".join(cells[:len(header)]) + " |")
            md_lines.append("\n")

        return "\n".join(md_lines)
    except Exception as e:
        return json.dumps({"error": str(e)}, indent=2)


@mcp.tool()
def get_document_outline(file_path: str) -> str:
    """
    Extract the heading hierarchy and table of contents tree (H1, H2, H3, H4) with paragraph positions.
    
    Args:
        file_path: Absolute or relative path to the .docx or .doc file.
    """
    try:
        doc = load_document(file_path)
        headings = []

        for p_idx, p in enumerate(doc.paragraphs):
            style_name = p.style.name if p.style else ""
            text = p.text.strip()
            if not text:
                continue

            level = None
            if style_name.startswith("Heading"):
                try:
                    level = int(style_name.replace("Heading", "").strip())
                except:
                    level = 1
            elif style_name == "Title":
                level = 0
            elif style_name == "Subtitle":
                level = 1

            if level is not None:
                headings.append({
                    "paragraph_index": p_idx,
                    "level": level,
                    "style": style_name,
                    "text": text
                })

        return json.dumps({
            "file_path": os.path.abspath(file_path),
            "heading_count": len(headings),
            "outline": headings
        }, indent=2)
    except Exception as e:
        return json.dumps({"error": str(e)}, indent=2)


@mcp.tool()
def get_document_metadata(file_path: str) -> str:
    """
    Retrieve document properties, author, title, revision, word count, and timestamp metadata.
    
    Args:
        file_path: Absolute or relative path to the .docx or .doc file.
    """
    try:
        doc = load_document(file_path)
        props = doc.core_properties

        total_paras = len(doc.paragraphs)
        total_words = sum(len(p.text.split()) for p in doc.paragraphs)
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    total_words += len(cell.text.split())

        meta = {
            "file_path": os.path.abspath(file_path),
            "title": props.title or "",
            "author": props.author or "",
            "subject": props.subject or "",
            "keywords": props.keywords or "",
            "comments": props.comments or "",
            "last_modified_by": props.last_modified_by or "",
            "revision": props.revision or 1,
            "created": props.created.isoformat() if props.created else None,
            "modified": props.modified.isoformat() if props.modified else None,
            "category": props.category or "",
            "statistics": {
                "total_paragraphs": total_paras,
                "total_tables": len(doc.tables),
                "total_sections": len(doc.sections),
                "estimated_word_count": total_words
            }
        }
        return json.dumps(meta, indent=2)
    except Exception as e:
        return json.dumps({"error": str(e)}, indent=2)


@mcp.tool()
def read_word_document(
    file_path: str,
    include_tables: bool = True,
    include_headers_footers: bool = True,
    output_format: str = "markdown",
    max_chars: int = 30000
) -> str:
    """
    Read and extract the content of a Word document formatted as Markdown or structured JSON.
    Includes token-safe truncation for large documents.
    
    Args:
        file_path: Path to the .docx or .doc file.
        include_tables: Whether to include table contents (default: True).
        include_headers_footers: Whether to include header and footer text (default: True).
        output_format: Output format ('markdown' or 'json'). Default is 'markdown'.
        max_chars: Maximum characters to return before truncating (default: 30000, ~7500 tokens).
    """
    try:
        doc = load_document(file_path)

        if output_format.lower() == "json":
            paragraphs_data = []
            for p_idx, p in enumerate(doc.paragraphs):
                if p.text.strip():
                    runs = []
                    for r in p.runs:
                        if r.text:
                            runs.append({
                                "text": r.text,
                                "formatting": extract_run_formatting(r, p.style)
                            })
                    paragraphs_data.append({
                        "index": p_idx,
                        "style": p.style.name if p.style else "",
                        "alignment": str(p.alignment).split(".")[-1] if p.alignment else "LEFT",
                        "text": p.text.strip(),
                        "runs": runs
                    })

            has_more = len(paragraphs_data) > 100
            display_paras = paragraphs_data[:100]
            result = {
                "file_path": os.path.abspath(file_path),
                "total_paragraphs": len(paragraphs_data),
                "returned_paragraphs": len(display_paras),
                "has_more": has_more,
                "paragraphs": display_paras,
            }
            if has_more:
                result["notice"] = "... [TRUNCATED: Showing first 100 paragraphs to protect context window] ..."
            if include_tables:
                tables_res = []
                for t in doc.tables[:20]:
                    rows = [[cell.text.strip() for cell in r.cells] for r in t.rows]
                    tables_res.append(rows)
                result["tables"] = tables_res
            return json.dumps(result, indent=2)

        md = []
        md.append(f"# Document: {os.path.basename(file_path)}\n")

        if include_headers_footers and doc.sections:
            sec = doc.sections[0]
            header_text = "\n".join(p.text.strip() for p in sec.header.paragraphs if p.text.strip())
            if header_text:
                md.append(f"> **Header**: {header_text}\n")

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style = p.style.name if p.style else "Normal"
            if style.startswith("Heading 1"):
                md.append(f"\n# {text}\n")
            elif style.startswith("Heading 2"):
                md.append(f"\n## {text}\n")
            elif style.startswith("Heading 3"):
                md.append(f"\n### {text}\n")
            elif style.startswith("Heading 4"):
                md.append(f"\n#### {text}\n")
            elif style == "Title":
                md.append(f"\n# {text}\n")
            elif style == "Subtitle":
                md.append(f"\n### *{text}*\n")
            elif "List" in style:
                md.append(f"- {text}")
            else:
                formatted_runs = []
                for r in p.runs:
                    r_text = r.text
                    if r.bold and r.italic:
                        r_text = f"***{r_text}***"
                    elif r.bold:
                        r_text = f"**{r_text}**"
                    elif r.italic:
                        r_text = f"*{r_text}*"
                    formatted_runs.append(r_text)
                md.append("".join(formatted_runs))

        if include_tables and doc.tables:
            md.append("\n## Embedded Tables\n")
            for t_idx, table in enumerate(doc.tables[:15], 1):
                md.append(f"### Table {t_idx}")
                rows = [[cell.text.strip().replace("\n", " ") for cell in r.cells] for r in table.rows]
                if rows:
                    header = rows[0]
                    md.append("| " + " | ".join(header) + " |")
                    md.append("| " + " | ".join(["---"] * len(header)) + " |")
                    for r in rows[1:50]:
                        cells = r + [""] * (len(header) - len(r))
                        md.append("| " + " | ".join(cells[:len(header)]) + " |")
                    if len(rows) > 50:
                        md.append(f"\n*... [{len(rows) - 50} table rows truncated] ...*\n")
                md.append("\n")

        if include_headers_footers and doc.sections:
            sec = doc.sections[0]
            footer_text = "\n".join(p.text.strip() for p in sec.footer.paragraphs if p.text.strip())
            if footer_text:
                md.append(f"\n> **Footer**: {footer_text}\n")

        full_output = "\n".join(md)
        if len(full_output) > max_chars:
            notice = f"\n\n... [TRUNCATED: Document exceeds {max_chars} characters. Use get_document_outline or search_word_document to inspect specific sections] ..."
            return full_output[:max_chars] + notice
        return full_output
    except Exception as e:
        return f"Error reading Word document: {str(e)}"


@mcp.tool()
def search_word_document(
    file_path: str,
    search_term: str,
    match_case: bool = False,
    max_matches: int = 50
) -> str:
    """
    Search for a text term or pattern across paragraphs, headers, footers, and table cells in a Word document.
    
    Args:
        file_path: Path to the .docx or .doc file.
        search_term: String or pattern to search for.
        match_case: Case-sensitive search flag (default: False).
        max_matches: Maximum number of matches to return (default: 50).
    """
    try:
        doc = load_document(file_path)
        matches = []
        flags = 0 if match_case else re.IGNORECASE
        pattern = re.compile(re.escape(search_term), flags)
        safe_max = min(max(1, max_matches), 200)

        for p_idx, p in enumerate(doc.paragraphs):
            if pattern.search(p.text):
                matches.append({
                    "location": f"Paragraph #{p_idx + 1}",
                    "style": p.style.name if p.style else "",
                    "text": p.text.strip()
                })
                if len(matches) >= safe_max:
                    break

        if len(matches) < safe_max:
            for s_idx, section in enumerate(doc.sections, 1):
                h_text = " ".join(p.text for p in section.header.paragraphs)
                if pattern.search(h_text):
                    matches.append({"location": f"Section {s_idx} Header", "text": h_text.strip()})
                f_text = " ".join(p.text for p in section.footer.paragraphs)
                if pattern.search(f_text):
                    matches.append({"location": f"Section {s_idx} Footer", "text": f_text.strip()})
                if len(matches) >= safe_max:
                    break

        if len(matches) < safe_max:
            for t_idx, table in enumerate(doc.tables, 1):
                for r_idx, row in enumerate(table.rows):
                    for c_idx, cell in enumerate(row.cells):
                        if pattern.search(cell.text):
                            matches.append({
                                "location": f"Table {t_idx}, Row {r_idx + 1}, Col {c_idx + 1}",
                                "text": cell.text.strip()
                            })
                            if len(matches) >= safe_max:
                                break
                    if len(matches) >= safe_max:
                        break
                if len(matches) >= safe_max:
                    break

        result_dict: Dict[str, Any] = {
            "file_path": os.path.abspath(file_path),
            "search_term": search_term,
            "total_matches": len(matches),
            "limit_reached": len(matches) >= safe_max,
            "matches": matches
        }
        if len(matches) >= safe_max:
            result_dict["notice"] = f"... [TRUNCATED: Showing first {safe_max} matches. Narrow your search term for more specific results] ..."

        return json.dumps(result_dict, indent=2)
    except Exception as e:
        return json.dumps({"error": str(e)}, indent=2)



if __name__ == "__main__":
    mcp.run()
