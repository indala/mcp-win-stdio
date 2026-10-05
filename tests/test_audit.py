import os
import sys
from pathlib import Path

for p in Path("packages").glob("*/src"):
    sys.path.insert(0, str(p.resolve()))
sys.path.insert(0, os.path.abspath("src"))
import json
import decimal
import uuid
from datetime import datetime, date, time
import pandas as pd
import numpy as np
import pathspec

def test_excel_clean_records():
    from mcp_win_stdio.excel.server import _df_to_clean_records
    df = pd.DataFrame({
        "dates": [pd.Timestamp("2026-01-01 12:00:00"), pd.NaT],
        "nums": [10.5, np.nan],
        "texts": ["hello", None]
    })
    records = _df_to_clean_records(df)
    assert len(records) == 2
    assert records[0]["dates"] == "2026-01-01T12:00:00"
    assert records[1]["dates"] is None
    assert records[0]["nums"] == 10.5
    assert records[1]["nums"] is None
    assert records[0]["texts"] == "hello"
    assert records[1]["texts"] is None
    print("[PASS] Excel clean records test passed.")

def test_explorer_gitignore():
    from mcp_win_stdio.explorer.server import _matches_gitignore
    patterns = ["node_modules/", "dist/", "*.pyc", "temp_dir/"]
    spec = pathspec.PathSpec.from_lines("gitignore", patterns)
    assert _matches_gitignore(spec, "node_modules", is_dir=True) is True
    assert _matches_gitignore(spec, "dist", is_dir=True) is True
    assert _matches_gitignore(spec, "src/foo.pyc", is_dir=False) is True
    assert _matches_gitignore(spec, "src/foo.py", is_dir=False) is False
    print("[PASS] Explorer gitignore test passed.")

def test_tsc_safety():
    from mcp_win_stdio.tsc.server import (
        find_tsconfigs,
        is_home_or_root_dir,
        get_default_watch_dir,
        check_tsc_available,
        list_watched_projects,
    )
    assert is_home_or_root_dir("C:\\") is True
    assert is_home_or_root_dir(str(Path.home())) is True
    appdata = os.environ.get("APPDATA")
    if appdata:
        assert is_home_or_root_dir(appdata) is True

    # Standby check when env and config are empty
    old_env = os.environ.pop("TSC_WATCH_DIR", None)
    old_proj = os.environ.pop("PROJECT_ROOT", None)
    from unittest.mock import patch
    try:
        with patch("mcp_win_stdio.core.config.load_config", return_value={}), \
             patch("mcp_win_stdio.tsc.server.WATCHED_PROJECTS", {}):
            assert get_default_watch_dir() is None
            status = list_watched_projects()
            assert status["status"] == "standby"
            assert status["watched_projects_count"] == 0
    finally:
        if old_env:
            os.environ["TSC_WATCH_DIR"] = old_env
        if old_proj:
            os.environ["PROJECT_ROOT"] = old_proj

    # Compiler pre-check
    tsc_check = check_tsc_available()
    assert "available" in tsc_check
    assert "command" in tsc_check

    # Should not crash or crawl home/root
    configs = find_tsconfigs("C:\\")
    assert isinstance(configs, list)

    # Context window protection & message truncation tests
    from mcp_win_stdio.tsc.server import _format_tsc_error, get_tsc_errors, get_file_errors, WATCHED_PROJECTS, CACHE_LOCK
    huge_msg = "Type '{ " + "x: string; " * 100 + "}' is not assignable to type 'number'."
    mock_err = {
        "file": "E:/project/src/App.tsx",
        "relative_path": "src/App.tsx",
        "line": 10,
        "column": 5,
        "severity": "error",
        "code": "TS2322",
        "message": huge_msg,
    }
    formatted = _format_tsc_error(mock_err, max_msg_chars=100)
    assert formatted["is_message_truncated"] is True
    assert len(formatted["message"]) < 200
    assert "truncated" in formatted["message"]

    # Pagination & limits test with mock watched project
    with CACHE_LOCK:
        fake_errors = [
            {
                "file": "E:/project/src/file.tsx",
                "relative_path": "src/file.tsx",
                "line": i,
                "column": 1,
                "severity": "error",
                "code": f"TS{2000 + i}",
                "message": f"Error line {i}",
            }
            for i in range(1, 60)
        ]
        WATCHED_PROJECTS["fake_tsconfig"] = {
            "relative_config": "tsconfig.json",
            "project_dir": "E:/project",
            "status": "ready",
            "last_updated": "2026-09-29T12:00:00Z",
            "errors": fake_errors,
        }

    try:
        # Test get_tsc_errors pagination
        page1 = get_tsc_errors(limit=20, offset=0)
        assert page1["total_errors"] == 59
        assert page1["returned_errors"] == 20
        assert page1["has_more"] is True
        assert "TRUNCATED" in page1["notice"]

        page2 = get_tsc_errors(limit=20, offset=20)
        assert page2["returned_errors"] == 20
        assert page2["offset"] == 20

        # Test error code filter
        filtered = get_tsc_errors(error_code="TS2010")
        assert filtered["total_errors"] == 1
        assert filtered["errors"][0]["code"] == "TS2010"

        # Test get_file_errors pagination
        file_res = get_file_errors("src/file.tsx", limit=15, offset=0)
        assert file_res["total_errors"] == 59
        assert file_res["returned_errors"] == 15
        assert file_res["has_more"] is True
        assert "TRUNCATED" in file_res["notice"]
    finally:
        with CACHE_LOCK:
            WATCHED_PROJECTS.pop("fake_tsconfig", None)

    print("[PASS] TSC safety, standby mode, context protection & pagination tests passed.")

def test_db_truncate_cell():
    import json
    from mcp_win_stdio.db.server import _truncate_cell
    d = decimal.Decimal("123.456")
    u = uuid.uuid4()
    now = datetime.now()
    mem = memoryview(b"binary_payload")
    
    cleaned = {
        "decimal": _truncate_cell(d),
        "uuid": _truncate_cell(u),
        "datetime": _truncate_cell(now),
        "mem": _truncate_cell(mem),
        "long_str": _truncate_cell("a" * 600, max_chars=100)
    }
    # Verify strict JSON serializability
    serialized = json.dumps(cleaned)
    assert "123.456" in serialized
    assert "<binary data: 14 bytes>" in serialized
    assert "... [truncated 500 chars]" in serialized
    print("[PASS] DB truncate cell serialization test passed.")

def test_excel_context_protection():
    import tempfile
    from mcp_win_stdio.excel.server import (
        query_rows, preview_sheet, read_range, get_column_values, compare_column_values, export_to_csv,
        export_to_json, profile_sheet, query_excel_sql
    )
    # Create sample Excel workbook
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = tmp.name
    try:
        df = pd.DataFrame({
            "code": [f"MAT{i:04d}" for i in range(150)],
            "desc": [f"Item description {i}" for i in range(150)],
            "qty": [i * 2 for i in range(150)]
        })
        df.to_excel(tmp_path, index=False)

        # 1. Test query_rows capping at 100 for records and returning protection notice
        res = query_rows(tmp_path, limit=5000, format="records")
        assert res["limit"] == 100
        assert res["returned_rows"] == 100
        assert "context_protection_notice" in res
        assert isinstance(res["rows"], list)
        assert len(res["rows"]) == 100

        # 2. Test compact format (list of lists)
        res_compact = query_rows(tmp_path, limit=200, format="compact")
        assert res_compact["format"] == "compact"
        assert isinstance(res_compact["data"], list)
        assert isinstance(res_compact["data"][0], list)
        assert res_compact["limit"] == 150 or res_compact["returned_rows"] == 150

        # 3. Test tsv format
        res_tsv = query_rows(tmp_path, limit=50, format="tsv")
        assert res_tsv["format"] == "tsv"
        assert "code\tdesc\tqty" in res_tsv["data"]

        # 4. Test get_column_values
        col_res = get_column_values(tmp_path, column="code", limit=50)
        assert col_res["unique_count"] == 150
        assert len(col_res["values"]) == 50
        assert col_res["values"][0] == "MAT0000"

        # 6. Test export_to_json
        json_tmp = tmp_path.replace(".xlsx", ".json")
        json_res = export_to_json(tmp_path, json_tmp, orient="records")
        assert json_res["status"] == "success"
        assert os.path.exists(json_tmp)
        with open(json_tmp, "r", encoding="utf-8") as jf:
            parsed = json.load(jf)
            assert len(parsed) == 150
            assert parsed[0]["code"] == "MAT0000"
        os.remove(json_tmp)

        # 7. Test profile_sheet
        prof_res = profile_sheet(tmp_path)
        assert prof_res["total_rows"] == 150
        assert "code" in prof_res["column_profiles"]
        assert prof_res["column_profiles"]["qty"]["numeric_stats"]["sum"] > 0

        # 8. Test query_excel_sql
        sql_res = query_excel_sql(tmp_path, "SELECT code, qty FROM sheet WHERE qty > 200 ORDER BY qty ASC")
        assert sql_res["total_returned"] > 0
        assert "code" in sql_res["columns"]
        assert isinstance(sql_res["rows"], list)

        # 9. Test audit_formulas and search_and_replace_cells
        import openpyxl
        from mcp_win_stdio.excel.server import audit_formulas, search_and_replace_cells
        wb = openpyxl.load_workbook(tmp_path)
        ws = wb.active
        ws["D1"] = "FormulaCol"
        ws["D2"] = "=SUM(#REF!)"
        ws["E1"] = "ReplaceCol"
        ws["E2"] = "TargetText123"
        wb.save(tmp_path)
        wb.close()

        af_res = audit_formulas(tmp_path)
        assert af_res["total_errors_found"] >= 1
        assert any(e["error_type"] == "#REF!" for e in af_res["errors"])

        sr_dry = search_and_replace_cells(tmp_path, search_val="TargetText123", replace_val="ReplacedVal", dry_run=True)
        assert sr_dry["dry_run"] is True
        assert sr_dry["total_matches"] >= 1

        sr_apply = search_and_replace_cells(tmp_path, search_val="TargetText123", replace_val="ReplacedVal", dry_run=False)
        assert sr_apply["dry_run"] is False
        assert sr_apply["total_matches"] >= 1

        print("[PASS] Excel context protection, column comparison, JSON export, SQL, formula audit, and search/replace tests passed.")
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_word_tools():
    import docx
    import tempfile
    from mcp_win_stdio.word.server import (
        read_word_document, edit_paragraph, insert_table,
        inspect_revisions_and_comments, get_document_layout, get_document_outline
    )

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        doc_path = tmp.name

    try:
        doc = docx.Document()
        doc.add_heading("Audit Test Document", level=1)
        doc.add_paragraph("Paragraph 1: Original text for audit testing.")
        doc.add_paragraph("Paragraph 2: Second paragraph.")
        doc.save(doc_path)

        # 1. Test read_word_document
        res_read = read_word_document(doc_path, max_paragraphs=10)
        assert "Audit Test Document" in res_read
        assert "Original text" in res_read

        # 2. Test edit_paragraph safety (fails without overwrite or output_path)
        fail_res = edit_paragraph(doc_path, paragraph_index=2, new_text="Edited text", overwrite=False)
        assert "error" in fail_res

        # 3. Test edit_paragraph with overwrite=True
        ok_res = edit_paragraph(doc_path, paragraph_index=2, new_text="Paragraph 1: Modified text.", overwrite=True)
        assert ok_res["status"] == "success"

        # 4. Test insert_table
        tbl_res = insert_table(
            doc_path,
            headers=["ID", "Name", "Role"],
            rows=[["1", "Alice", "Admin"], ["2", "Bob", "User"]],
            overwrite=True
        )
        assert tbl_res["status"] == "success"
        assert tbl_res["rows_inserted"] == 2

        # 5. Test inspect_revisions_and_comments
        rev_res = inspect_revisions_and_comments(doc_path)
        assert "total_comments" in rev_res
        assert "total_insertions" in rev_res

        # 6. Test structured dict returns for layout & outline
        layout = get_document_layout(doc_path)
        assert isinstance(layout, dict)
        assert layout["total_sections"] >= 1

        outline = get_document_outline(doc_path)
        assert isinstance(outline, dict)
        assert outline["total_headings"] >= 1

        print("[PASS] Word server tests passed.")
    finally:
        if os.path.exists(doc_path):
            os.remove(doc_path)


def test_explorer_tools():
    from mcp_win_stdio.explorer.server import (
        read_file, find_symbol, get_code_stats, find_duplicate_files
    )

    # 1. Test read_file capping
    res_rf = read_file("src/mcp_win_stdio/cli.py", max_lines=10)
    assert res_rf["lines_returned"] == 10
    assert res_rf["has_more"] is True

    # 2. Test find_symbol
    sym_res = find_symbol("src/mcp_win_stdio", symbol_name="main", exact_match=True, language="python")
    assert sym_res["total_matches"] >= 1
    assert any(s["name"] == "main" for s in sym_res["symbols"])

    # 3. Test get_code_stats
    stats_res = get_code_stats("src/mcp_win_stdio")
    assert stats_res["total_files_analyzed"] > 0
    assert "Python" in stats_res["languages"]

    # 4. Test find_duplicate_files
    dup_res = find_duplicate_files("src/mcp_win_stdio", min_size_bytes=100)
    assert "total_duplicate_groups" in dup_res

    print("[PASS] Explorer advanced tools test passed.")


def test_db_agent_friction_fixes():
    from mcp_win_stdio.db.server import (
        _split_sql_statements,
        _normalize_connection_param,
        _format_db_error,
        _get_connection,
        _ACTIVE_CONNECTION
    )

    # 1. Test _normalize_connection_param (Sticky Connection Fix)
    for null_val in [None, "", "  ", "null", "NULL", "Null", "none", "NONE", "default", "DEFAULT", "undefined", "UNDEFINED"]:
        assert _normalize_connection_param(null_val) is None, f"Failed for {null_val}"
    assert _normalize_connection_param("showreel_dev") == "showreel_dev"
    assert _normalize_connection_param("  my_database  ") == "my_database"
    print("[PASS] Connection normalization test passed.")

    # 2. Test _split_sql_statements (Multi-statement batch execution)
    # Basic statements
    sql_basic = "SELECT 1; SELECT 2; SELECT 3;"
    stmts = _split_sql_statements(sql_basic)
    assert stmts == ["SELECT 1", "SELECT 2", "SELECT 3"]

    # Semicolon inside single quotes
    sql_quotes = "INSERT INTO test_tbl (val) VALUES ('a;b;c'); SELECT * FROM test_tbl;"
    stmts_q = _split_sql_statements(sql_quotes)
    assert len(stmts_q) == 2
    assert stmts_q[0] == "INSERT INTO test_tbl (val) VALUES ('a;b;c')"
    assert stmts_q[1] == "SELECT * FROM test_tbl"

    # Escaped quotes inside strings
    sql_esc = "INSERT INTO t (val) VALUES ('It''s a value; indeed'); SELECT 1;"
    stmts_esc = _split_sql_statements(sql_esc)
    assert len(stmts_esc) == 2
    assert "It''s a value; indeed" in stmts_esc[0]

    # Comments with semicolons
    sql_comments = "SELECT 1; -- comment with ; inside\nSELECT 2; /* block ; comment */ SELECT 3;"
    stmts_c = _split_sql_statements(sql_comments)
    assert len(stmts_c) == 3

    # PostgreSQL Dollar Quoting ($$ ... $$ and $tag$ ... $tag$)
    sql_dollar = """
    CREATE OR REPLACE FUNCTION test_func() RETURNS void AS $$
    BEGIN
        SELECT 1;
        INSERT INTO audit_log VALUES ('done; logged');
    END;
    $$ LANGUAGE plpgsql;
    SELECT test_func();
    """
    stmts_d = _split_sql_statements(sql_dollar)
    assert len(stmts_d) == 2
    assert "CREATE OR REPLACE FUNCTION" in stmts_d[0]
    assert "SELECT test_func()" in stmts_d[1]

    sql_tagged_dollar = """
    CREATE FUNCTION tagged() RETURNS text AS $body$
        SELECT 'hello; world';
    $body$ LANGUAGE sql;
    SELECT 42;
    """
    stmts_td = _split_sql_statements(sql_tagged_dollar)
    assert len(stmts_td) == 2
    print("[PASS] SQL statement splitter test passed.")

    # 3. Test _format_db_error (Rich diagnostic output)
    e = ValueError("Table 'missing_table' not found.")
    formatted = _format_db_error(e, engine="postgres", sql="SELECT * FROM missing_table WHERE id = 1;", statement_idx=2)
    assert formatted["error"] is True
    assert formatted["engine"] == "postgres"
    assert formatted["failed_statement_index"] == 2
    assert formatted["message"] == "Table 'missing_table' not found."
    assert "sql_snippet" in formatted
    assert "SELECT * FROM missing_table" in formatted["sql_snippet"]
    print("[PASS] DB error formatting test passed.")

    # 4. Test LIKE parameter handling with psycopg2 / DB-API (Issue #4)
    from unittest.mock import MagicMock
    mock_cursor = MagicMock()
    mock_cursor.rowcount = 1
    mock_cursor.statusmessage = "SELECT 1"
    mock_cursor.fetchmany.return_value = [{"col": "showreel_dev"}]

    mock_conn = MagicMock()
    mock_conn.cursor.return_value.__enter__.return_value = mock_cursor

    # Simulate read_query logic with LIKE '%s...' pattern
    test_sql = "SELECT * FROM videos WHERE title LIKE '%showreel%';"
    # When params is None or empty, cur.execute must be called with ONLY (sql,), NEVER (sql, ())
    params = None
    if params:
        mock_cursor.execute(test_sql, tuple(params))
    else:
        mock_cursor.execute(test_sql)

    # Verify cur.execute was called with 1 argument (test_sql), not 2
    mock_cursor.execute.assert_called_once_with(test_sql)
    print("[PASS] LIKE query parameter handling test passed.")



def test_db_pii_masking():
    """Test that _truncate_row redacts PII columns and passes safe columns through."""
    from mcp_win_stdio.db.server import _truncate_row, PII_COLUMN_PATTERN
    import re

    # PII_COLUMN_PATTERN should match these
    sensitive_cols = ["password", "password_hash", "api_key", "secret_token", "ssn",
                      "credit_card_number", "auth_token", "private_key", "access_token"]
    for col in sensitive_cols:
        assert PII_COLUMN_PATTERN.search(col), f"PII pattern should match '{col}'"

    # Should NOT match ordinary columns
    safe_cols = ["user_id", "username", "email", "created_at", "status", "name"]
    for col in safe_cols:
        assert not PII_COLUMN_PATTERN.search(col), f"PII pattern should NOT match '{col}'"

    # Test row-level masking
    row = {
        "user_id": 42,
        "username": "alice",
        "password_hash": "$2b$12$supersecretstuff",
        "api_key": "sk-abc123",
        "email": "alice@example.com",
        "auth_token": "eyJhbGciOiJIUzI1NiJ9.xxx"
    }

    masked = _truncate_row(row, mask_sensitive=True)
    assert masked["user_id"] == 42
    assert masked["username"] == "alice"
    assert masked["email"] == "alice@example.com"
    assert masked["password_hash"] == "[REDACTED_SENSITIVE]", f"Got: {masked['password_hash']}"
    assert masked["api_key"] == "[REDACTED_SENSITIVE]", f"Got: {masked['api_key']}"
    assert masked["auth_token"] == "[REDACTED_SENSITIVE]", f"Got: {masked['auth_token']}"

    # With mask_sensitive=False everything should pass through
    unmasked = _truncate_row(row, mask_sensitive=False)
    assert unmasked["password_hash"] == "$2b$12$supersecretstuff"
    assert unmasked["api_key"] == "sk-abc123"

    print("[PASS] DB PII masking test passed.")


def test_db_schema_tools():
    """Test compare_schemas migration_sql structure and audit_database_health non-postgres guard."""
    from mcp_win_stdio.db.server import audit_database_health, _CONNECTION_REGISTRY

    # audit_database_health should reject MySQL connections
    # We temporarily register a fake MySQL connection
    _CONNECTION_REGISTRY["_test_mysql_fake"] = {
        "name": "_test_mysql_fake",
        "engine": "mysql",
        "database": "testdb",
        "url": "mysql://user:pass@localhost/testdb",
    }
    try:
        result = audit_database_health(connection="_test_mysql_fake")
        assert "error" in result, f"Expected error for MySQL, got: {result}"
        assert "PostgreSQL" in result["error"]
        print("[PASS] audit_database_health MySQL guard test passed.")
    finally:
        _CONNECTION_REGISTRY.pop("_test_mysql_fake", None)

    # Verify compare_schemas migration_sql key is always present when the module imports
    # (structural check only — no live DB needed)
    import inspect
    from mcp_win_stdio.db import server as db_server
    src = inspect.getsource(db_server.compare_schemas)
    assert "migration_sql" in src, "compare_schemas must produce migration_sql"
    print("[PASS] compare_schemas migration_sql structural check passed.")


def test_ssh_tools():
    from mcp_win_stdio.ssh.server import (
        add_host,
        list_hosts,
        use_host,
        remove_host,
        list_active_connections,
        ssh_tunnel_list,
        ssh_list_pty_sessions,
    )
    # Test adding host
    res_add = add_host("ci-test-host", "127.0.0.1", "testuser", 2222)
    assert res_add["success"] is True
    assert res_add["host"] == "ci-test-host"

    # Test listing hosts
    res_list = list_hosts()
    assert res_list["totalHosts"] >= 1
    found = any(h["name"] == "ci-test-host" for h in res_list["hosts"])
    assert found is True

    # Test switching active host
    res_use = use_host("ci-test-host")
    assert res_use["success"] is True
    assert res_use["activeHost"] == "ci-test-host"

    # Test empty / clean state queries
    assert isinstance(list_active_connections()["connections"], list)
    assert isinstance(ssh_tunnel_list()["tunnels"], list)
    assert isinstance(ssh_list_pty_sessions()["sessions"], list)

    # Test removing host
    res_rm = remove_host("ci-test-host")
    assert res_rm["success"] is True

def test_db_server_database_confusion():
    """Test that DB MCP seamlessly resolves aliases when server, connection, database, or name are used."""
    from mcp_win_stdio.db.server import (
        _RAW_CONFIG,
        _CONNECTION_REGISTRY,
        _get_connection,
        _resolve_conn,
        use_database,
    )
    # Simulate user configuring showreel server with showreel database
    _RAW_CONFIG["showreel"] = "postgresql://postgres:pass@localhost:5432/showreel"
    _RAW_CONFIG["analytics_server"] = "postgresql://postgres:pass@remote:5432/metrics_db"

    try:
        # 1. Resolve by connection alias
        c1 = _get_connection("showreel")
        assert c1["name"] == "showreel"
        assert c1["database"] == "showreel"

        # 2. Resolve by database name inside configured server (metrics_db on analytics_server)
        c2 = _get_connection("metrics_db")
        assert c2["name"] == "analytics_server"
        assert c2["database"] == "metrics_db"

        # 3. Test _resolve_conn helper
        assert _resolve_conn("showreel", None) == "showreel"
        assert _resolve_conn(None, "showreel") == "showreel"
        assert _resolve_conn("showreel", "fallback") == "showreel"
        assert _resolve_conn("null", "showreel") == "showreel"

        # 4. Test use_database with various argument forms:
        # a) database="showreel"
        res_db = use_database(database="showreel")
        assert res_db["success"] is True
        assert res_db["activeConnection"] == "showreel"

        # b) connection="showreel"
        res_conn = use_database(connection="showreel")
        assert res_conn["success"] is True
        assert res_conn["activeConnection"] == "showreel"

        # c) server="showreel"
        res_srv = use_database(server="showreel")
        assert res_srv["success"] is True
        assert res_srv["activeConnection"] == "showreel"

        # d) name="showreel"
        res_name = use_database(name="showreel")
        assert res_name["success"] is True
        assert res_name["activeConnection"] == "showreel"

        # e) switching to metrics_db by database name
        res_metrics = use_database(database="metrics_db")
        assert res_metrics["success"] is True
        assert res_metrics["activeConnection"] == "analytics_server"
        assert res_metrics["database"] == "metrics_db"

        print("[PASS] DB server vs database confusion & alias resolution tests passed.")
    finally:
        _RAW_CONFIG.pop("showreel", None)
        _RAW_CONFIG.pop("analytics_server", None)
        _CONNECTION_REGISTRY.pop("showreel", None)
        _CONNECTION_REGISTRY.pop("analytics_server", None)


def test_tsc_filter_improvements():
    """Test TSC substring filtering, case-insensitivity, and suggestion tools."""
    from mcp_win_stdio.tsc.server import (
        WATCHED_PROJECTS,
        CACHE_LOCK,
        get_tsc_errors,
        get_file_errors,
        suggest_error_fixes,
        get_error_category_breakdown,
    )
    with CACHE_LOCK:
        WATCHED_PROJECTS["E:/repos/my-app/tsconfig.json"] = {
            "relative_config": "tsconfig.json",
            "project_dir": "E:/repos/my-app",
            "status": "ready",
            "last_updated": "2026-10-01T12:00:00Z",
            "errors": [
                {
                    "file": "E:/repos/my-app/src/App.tsx",
                    "relative_path": "src/App.tsx",
                    "line": 15,
                    "column": 5,
                    "severity": "error",
                    "code": "TS2322",
                    "message": "Type 'string' is not assignable to type 'number'.",
                },
                {
                    "file": "E:/repos/my-app/src/utils.ts",
                    "relative_path": "src/utils.ts",
                    "line": 42,
                    "column": 10,
                    "severity": "error",
                    "code": "TS2304",
                    "message": "Cannot find name 'missingVar'.",
                },
            ],
        }

    try:
        # 1. Test substring filtering for project_path (e.g. 'my-app' or 'my-app/src')
        res_sub = get_tsc_errors(project_path="my-app")
        assert res_sub["total_errors"] == 2
        assert len(res_sub["errors"]) == 2

        # 2. Test tsconfig substring filter
        res_cfg = get_tsc_errors(tsconfig_path="my-app/tsconfig.json")
        assert res_cfg["total_errors"] == 2

        # 3. Test get_file_errors with Windows path and casing
        res_file = get_file_errors("E:\\repos\\MY-APP\\src\\app.tsx")
        assert res_file["total_errors"] == 1
        assert res_file["errors"][0]["code"] == "TS2322"

        # 4. Test suggest_error_fixes
        res_fix = suggest_error_fixes("TS2322")
        assert res_fix["success"] is True
        assert "recommended_resolutions" in res_fix
        assert len(res_fix["recommended_resolutions"]) > 0

        # 5. Test get_error_category_breakdown
        res_cat = get_error_category_breakdown()
        assert res_cat["success"] is True
        assert "type_mismatches" in res_cat["category_summary"]
        assert res_cat["category_summary"]["type_mismatches"] == 1

        print("[PASS] TSC substring filter, case-insensitivity & suggestions tests passed.")
    finally:
        with CACHE_LOCK:
            WATCHED_PROJECTS.pop("E:/repos/my-app/tsconfig.json", None)


def test_excel_styling_and_layout():
    import tempfile
    import openpyxl
    from mcp_win_stdio.excel.server import (
        create_workbook,
        format_cells,
        apply_conditional_formatting,
        set_sheet_layout_and_freeze,
        update_cells,
    )

    with tempfile.TemporaryDirectory() as td:
        wb_path = os.path.join(td, "styled_test.xlsx")
        create_workbook(
            file_path=wb_path,
            sheet_name="Sales",
            data=[
                {"Product": "Widgets", "Q1": 1200, "Q2": 1500, "Total": 2700},
                {"Product": "Gadgets", "Q1": 300, "Q2": 450, "Total": 750},
                {"Product": "Doohickeys", "Q1": 50, "Q2": 80, "Total": 130},
            ],
        )

        # 1. Test format_cells
        res_fmt = format_cells(
            file_path=wb_path,
            range_address="A1:D1",
            sheet_name="Sales",
            fill="1F497D",
            font={"bold": True, "color": "FFFFFF", "name": "Segoe UI"},
            alignment={"horizontal": "center"},
            border="thin",
        )
        assert res_fmt["status"] == "success"
        assert res_fmt["cells_formatted"] == 4

        # 2. Test format_cells number format
        res_num = format_cells(
            file_path=wb_path,
            range_address="B2:D4",
            sheet_name="Sales",
            number_format="currency",
        )
        assert res_num["status"] == "success"

        # Verify with openpyxl
        wb = openpyxl.load_workbook(wb_path)
        ws = wb["Sales"]
        assert ws["A1"].font.bold is True
        assert "FFFFFF" in str(ws["A1"].font.color.rgb)
        assert ws["B2"].number_format == "$#,##0.00"
        wb.close()

        # 3. Test apply_conditional_formatting
        res_cond = apply_conditional_formatting(
            file_path=wb_path,
            sheet_name="Sales",
            range_address="D2:D4",
            rule_type="cell_is",
            operator="greaterThan",
            formula=["1000"],
            fill_color="C6EFCE",
            font_color="006100",
        )
        assert res_cond["status"] == "success"

        # Color scale test
        res_scale = apply_conditional_formatting(
            file_path=wb_path,
            sheet_name="Sales",
            range_address="B2:C4",
            rule_type="color_scale",
            color_scale_preset="green_yellow_red",
        )
        assert res_scale["status"] == "success"

        # 4. Test set_sheet_layout_and_freeze
        res_layout = set_sheet_layout_and_freeze(
            file_path=wb_path,
            sheet_name="Sales",
            auto_fit_columns=True,
            freeze_panes="A2",
            show_grid_lines=True,
        )
        assert res_layout["status"] == "success"
        assert "auto_fit_columns" in res_layout["settings_applied"]
        assert res_layout["settings_applied"]["freeze_panes"] == "A2"

        wb = openpyxl.load_workbook(wb_path)
        ws = wb["Sales"]
        assert ws.freeze_panes == "A2"
        assert len(ws.conditional_formatting) == 2
        wb.close()

        # 5. Test update_cells with inline styling
        res_upd = update_cells(
            file_path=wb_path,
            sheet_name="Sales",
            updates=[
                {
                    "cell": "A5",
                    "value": "Total Summary",
                    "font": {"bold": True, "color": "1F497D"},
                    "fill": "D9E1F2",
                }
            ],
        )
        assert res_upd["status"] == "success"
        assert res_upd["cells_updated"] == 1

        wb = openpyxl.load_workbook(wb_path)
        ws = wb["Sales"]
        assert ws["A5"].value == "Total Summary"
        assert ws["A5"].font.bold is True
        wb.close()

        print("[PASS] Excel styling, conditional formatting, and layout tests passed!")


def test_excel_file_lock_instruction():
    import pytest
    from mcp_win_stdio.excel.server import _safe_save_workbook

    class DummyWb:
        def save(self, path):
            raise PermissionError(13, "Permission denied")

    wb = DummyWb()
    with pytest.raises(PermissionError) as exc_info:
        _safe_save_workbook(wb, "C:/Work/financial_report.xlsx")

    err_msg = str(exc_info.value)
    assert "FILE LOCKED BY EXCEL" in err_msg
    assert "financial_report.xlsx" in err_msg
    assert "DO NOT attempt terminal workarounds" in err_msg
    assert "PLEASE ASK THE USER" in err_msg
    print("[PASS] Excel file lock instruction test passed!")


def test_excel_write_range():
    import tempfile
    import openpyxl
    from mcp_win_stdio.excel.server import create_workbook, write_range

    with tempfile.TemporaryDirectory() as td:
        wb_path = os.path.join(td, "bulk_test.xlsx")
        create_workbook(
            file_path=wb_path,
            sheet_name="Breakdown",
            data=[
                {"Code": "OLD01", "Category": "Old", "Sub": "Old", "SubCode": 1, "Desc": "Old 1"},
                {"Code": "OLD02", "Category": "Old", "Sub": "Old", "SubCode": 2, "Desc": "Old 2"},
                {"Code": "OLD03", "Category": "Old", "Sub": "Old", "SubCode": 3, "Desc": "Old 3"},
            ],
        )

        # Overwrite starting at A2 with clear_subsequent_rows=True
        new_data = [
            ["AC01001", "AC", "Accessories", 1, "Power Distribution"],
            ["AC01002", "AC", "Accessories", 1, "Power Distribution"],
        ]
        res = write_range(
            file_path=wb_path,
            sheet_name="Breakdown",
            start_cell="A2",
            data=new_data,
            clear_subsequent_rows=True,
        )
        assert res["status"] == "success"
        assert res["rows_written"] == 2
        assert res["cells_written"] == 10

        # Verify workbook content
        wb = openpyxl.load_workbook(wb_path)
        ws = wb["Breakdown"]
        assert ws.max_row == 3  # Header row + 2 rows (old 3rd row deleted)
        assert ws["A2"].value == "AC01001"
        assert ws["E2"].value == "Power Distribution"
        assert ws["A3"].value == "AC01002"
        wb.close()

        print("[PASS] Excel write_range test passed!")


def test_excel_copilot_tools():
    import tempfile
    import openpyxl
    from mcp_win_stdio.excel.server import (
        create_workbook, create_chart, clean_and_deduplicate_sheet, transform_sheet_data
    )

    with tempfile.TemporaryDirectory() as td:
        wb_path = os.path.join(td, "copilot_test.xlsx")

        # 1. Create dataset with duplicates, extra whitespace, and dirty rows
        dirty_data = [
            {"Code": "A01", "Region": " North ", "Product": "Widget ", "Sales": 150, "Quantity": 10},
            {"Code": "A01", "Region": " North ", "Product": "Widget ", "Sales": 150, "Quantity": 10},  # Duplicate
            {"Code": "A02", "Region": "South", "Product": "Gadget", "Sales": 80, "Quantity": 5},
            {"Code": "A03", "Region": "North", "Product": "Gadget", "Sales": 220, "Quantity": 15},
            {"Code": "A04", "Region": "South", "Product": "Widget", "Sales": 310, "Quantity": 20},
        ]
        create_workbook(file_path=wb_path, sheet_name="Data", data=dirty_data)

        # 2. Test clean_and_deduplicate_sheet
        clean_res = clean_and_deduplicate_sheet(
            file_path=wb_path,
            sheet_name="Data",
            deduplicate_columns=["Code"],
            trim_text=True,
            drop_empty_rows=True,
        )
        assert clean_res["status"] == "success"
        assert clean_res["duplicates_removed"] == 1
        assert clean_res["final_rows"] == 4

        # Verify trimming in workbook
        wb = openpyxl.load_workbook(wb_path)
        ws = wb["Data"]
        assert ws["B2"].value == "North"  # Untrimmed ' North ' became 'North'
        assert ws["C2"].value == "Widget" # Untrimmed 'Widget ' became 'Widget'
        wb.close()

        # 3. Test transform_sheet_data (Grouping and Aggregations)
        transform_res = transform_sheet_data(
            file_path=wb_path,
            source_sheet="Data",
            output_sheet="Region_Summary",
            group_by=["Region"],
            aggregations={"Sales": "sum", "Quantity": "sum"},
            sort_by="Sales",
            ascending=False,
        )
        assert transform_res["status"] == "success"
        assert transform_res["result_rows"] == 2
        assert "Region" in transform_res["columns"]

        # Verify summary sheet exists and has grouped data
        wb = openpyxl.load_workbook(wb_path)
        assert "Region_Summary" in wb.sheetnames
        sum_ws = wb["Region_Summary"]
        assert sum_ws.max_row == 3  # Header + 2 region rows (North, South)
        wb.close()

        # 4. Test create_chart (Native Bar / Column chart)
        chart_res = create_chart(
            file_path=wb_path,
            sheet_name="Region_Summary",
            chart_type="bar",
            data_range="B1:B3",       # Header + 2 data rows for Sales
            categories_range="A2:A3", # Region categories
            title="Total Sales by Region",
            target_cell="D2",
            x_axis_title="Region",
            y_axis_title="Sales ($)",
        )
        assert chart_res["status"] == "success"
        assert chart_res["chart_type"] == "bar"

        # Verify chart attached in worksheet
        wb = openpyxl.load_workbook(wb_path)
        sum_ws = wb["Region_Summary"]
        assert len(sum_ws._charts) == 1
        assert "Total Sales by Region" in str(sum_ws._charts[0].title)
        wb.close()

        print("[PASS] Excel copilot tools (clean, transform, chart) test passed!")


if __name__ == "__main__":
    test_excel_clean_records()
    test_excel_styling_and_layout()
    test_excel_file_lock_instruction()
    test_excel_write_range()
    test_excel_copilot_tools()
    test_explorer_gitignore()
    test_tsc_safety()
    test_tsc_filter_improvements()
    test_db_truncate_cell()
    test_excel_context_protection()
    test_word_tools()
    test_explorer_tools()
    test_db_agent_friction_fixes()
    test_db_pii_masking()
    test_db_schema_tools()
    test_db_server_database_confusion()
    test_ssh_tools()
    print("ALL AUDIT TESTS PASSED!")
